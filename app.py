import os
import re
import time
import glob
import json
import base64
import hashlib
import html
import sqlite3
import platform
import smtplib
import ssl
import uuid
import urllib.parse
from email.message import EmailMessage
from datetime import datetime, date, time as dt_time
from typing import List, Dict, Optional, Tuple

import requests

import pandas as pd
import streamlit as st
import soundfile as sf
from openai import OpenAI
from docx import Document

try:
    import sounddevice as sd
    SOUNDDEVICE_AVAILABLE = True
except Exception:
    sd = None
    SOUNDDEVICE_AVAILABLE = False

from whisper_local import transcribe_whisper
from prompts import SYSTEM_INSTRUCTIONS
from app.hl7.audit import audit_event
from app.hl7.config import HL7Settings
from app.hl7.escape import compute_payload_sha256, prevent_phi_in_logs
from app.hl7.messages import build_adt_a04 as build_adt_a04_message
from app.hl7.messages import build_mdm_t02 as build_mdm_t02_message
from app.hl7.storage import enqueue_message, ensure_hl7_tables, get_latest_message, get_message, update_message_status
from app.hl7.transport import FileHL7Transport, get_hl7_transport
from app.hl7.validators import HL7ValidationError, get_control_id, parse_ack, validate_message_ready
from app.security import is_insecure_default_password, validate_admin_password_policy

try:
    from passlib.context import CryptContext
    PASSWORD_CONTEXT = CryptContext(schemes=["argon2"], deprecated="auto")
except Exception:
    PASSWORD_CONTEXT = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS_DIR = os.path.join(BASE_DIR, "assets")
APP_ICON_PATH = os.path.join(ASSETS_DIR, "asistente_simon_icon.png")
APP_NAME = "Simon Assistant"


def get_setting(name: str, default: str = "") -> str:
    value = os.getenv(name)
    if value not in (None, ""):
        return str(value)
    try:
        value = st.secrets.get(name, None)
    except Exception:
        value = None
    if value in (None, ""):
        return default
    return str(value)


def get_bool_setting(name: str, default: bool = False) -> bool:
    raw = get_setting(name, "")
    if raw == "":
        return default
    return raw.strip().lower() in {"1", "true", "yes", "si", "sí", "on"}


def get_int_setting(name: str, default: int) -> int:
    raw = get_setting(name, "")
    if raw == "":
        return default
    try:
        return int(raw)
    except ValueError:
        return default

# =========================================================
# Proyecto Jorge Ambulatorio - Windows
# Ficha clinica ambulatorio + IA + CIE-10 + HL7 v2 preparado
# =========================================================

OPENAI_API_KEY = get_setting("OPENAI_API_KEY", "")


def get_openai_client() -> OpenAI:
    """Crea el cliente OpenAI solo si existe una API key válida en Secrets/entorno."""
    api_key = get_setting("OPENAI_API_KEY", "")
    if not api_key:
        st.error(
            "Falta configurar OPENAI_API_KEY en los Secrets de Streamlit Cloud. "
            "No pongas la API key dentro del código ni en GitHub."
        )
        st.stop()
    return OpenAI(api_key=api_key)
os.makedirs(ASSETS_DIR, exist_ok=True)

ADMIN_ROLE = "admin"
DOCTOR_ROLE = "doctor"
DEFAULT_ADMIN_EMAIL = get_setting("PJ_ADMIN_EMAIL", "kinepdiaz@gmail.com").strip().lower()
APP_ENV = get_setting("APP_ENV", "development").strip().lower() or "development"
DEFAULT_ADMIN_PASSWORD = get_setting("PJ_ADMIN_PASSWORD", "")
LEGACY_DEFAULT_ADMIN_EMAIL = "admin@gmail.com"


def normalize_user_role(role: str) -> str:
    return ADMIN_ROLE if role == ADMIN_ROLE else DOCTOR_ROLE


def is_admin_user(user: Optional[dict]) -> bool:
    return bool(user and user.get("role") == ADMIN_ROLE)


def get_page_icon():
    if not os.path.exists(APP_ICON_PATH):
        return None
    try:
        from PIL import Image
        return Image.open(APP_ICON_PATH)
    except Exception:
        return APP_ICON_PATH


def inject_custom_css():
    st.markdown(
        '''
        <style>
        :root {
            --as-blue: #0F4C81;
            --as-blue-hover: #123E66;
            --as-teal: #1B998B;
            --as-teal-hover: #147A70;
            --as-bg: #F6F8FB;
            --as-sidebar: #EEF4F8;
            --as-card: #FFFFFF;
            --as-soft-blue: #D9EEF7;
            --as-text: #172033;
            --as-muted: #64748B;
            --as-border: #D8E2EC;
            --as-success: #15803D;
            --as-warning: #B45309;
            --as-error: #B91C1C;
            --as-radius: 16px;
            --as-shadow: 0 8px 24px rgba(15, 76, 129, 0.08);
            --pj-bg: var(--as-bg);
            --pj-surface: var(--as-card);
            --pj-surface-2: var(--as-sidebar);
            --pj-primary: var(--as-blue);
            --pj-primary-hover: var(--as-blue-hover);
            --pj-secondary: var(--as-teal);
            --pj-accent-soft: var(--as-soft-blue);
            --pj-text: var(--as-text);
            --pj-muted: var(--as-muted);
            --pj-border: var(--as-border);
            --pj-success: var(--as-success);
            --pj-warning: var(--as-warning);
            --pj-error: var(--as-error);
            --pj-radius: var(--as-radius);
            --pj-shadow: var(--as-shadow);
        }

        .stApp {
            background: var(--as-bg);
            color: var(--as-text);
        }

        .block-container {
            padding-top: 2rem;
            padding-bottom: 3rem;
            max-width: 1320px;
        }

        section[data-testid="stSidebar"] {
            background: linear-gradient(180deg, #FFFFFF 0%, var(--as-sidebar) 100%);
            border-right: 1px solid var(--as-border);
        }

        section[data-testid="stSidebar"] h1,
        section[data-testid="stSidebar"] h2,
        section[data-testid="stSidebar"] h3 {
            color: var(--as-blue);
            font-weight: 800;
        }

        section[data-testid="stSidebar"] .stButton > button {
            width: 100%;
            justify-content: flex-start;
        }

        section[data-testid="stSidebar"] .stButton > button[kind="secondary"] {
            background: rgba(255,255,255,0.72);
            border: 1px solid transparent;
            border-left: 5px solid transparent;
            color: var(--as-text);
            box-shadow: none;
            font-weight: 700;
        }

        section[data-testid="stSidebar"] .stButton > button[kind="secondary"]:hover {
            background: #FFFFFF;
            border-color: rgba(15, 76, 129, 0.16);
            border-left-color: var(--as-soft-blue);
            color: var(--as-blue);
            transform: translateY(-1px);
            box-shadow: 0 6px 14px rgba(15, 76, 129, 0.08);
        }

        section[data-testid="stSidebar"] .stButton > button[kind="primary"] {
            background: var(--as-blue);
            border: 1px solid var(--as-blue);
            border-left: 5px solid var(--as-teal);
            color: #FFFFFF;
            box-shadow: 0 8px 18px rgba(15, 76, 129, 0.18);
            font-weight: 800;
        }

        section[data-testid="stSidebar"] .stButton > button[kind="primary"]:hover {
            background: var(--as-blue-hover);
            border-color: var(--as-blue-hover);
            border-left-color: var(--as-teal);
            color: #FFFFFF;
        }

        h1, h2, h3 {
            color: var(--as-text);
            letter-spacing: 0;
        }

        p, label, span, div {
            font-family: system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
        }

        div[data-testid="stMetric"] {
            background: var(--as-card);
            border: 1px solid var(--as-border);
            border-radius: var(--as-radius);
            padding: 1rem 1.15rem;
            box-shadow: var(--as-shadow);
        }

        div[data-testid="stMetric"] label {
            color: var(--as-muted) !important;
        }

        div[data-testid="stExpander"] {
            background: var(--as-card);
            border: 1px solid var(--as-border);
            border-radius: var(--as-radius);
            box-shadow: 0 4px 16px rgba(15, 76, 129, 0.05);
            overflow: hidden;
            margin-bottom: 1rem;
        }

        div[data-testid="stExpander"] summary {
            font-weight: 800;
            color: var(--as-blue);
        }

        div[data-testid="stForm"],
        div[data-testid="stVerticalBlockBorderWrapper"] {
            border-color: var(--as-border);
            border-radius: var(--as-radius);
            box-shadow: 0 4px 16px rgba(15, 76, 129, 0.04);
        }

        .stButton > button,
        .stFormSubmitButton > button {
            border-radius: 999px;
            border: 1px solid var(--as-blue);
            background: var(--as-blue);
            color: #FFFFFF;
            font-weight: 700;
            min-height: 2.55rem;
            padding: 0.55rem 1.05rem;
            transition: all 0.15s ease-in-out;
        }

        .stButton > button:hover,
        .stFormSubmitButton > button:hover {
            background: var(--as-blue-hover);
            border-color: var(--as-blue-hover);
            color: #FFFFFF;
            transform: translateY(-1px);
            box-shadow: 0 8px 18px rgba(15, 76, 129, 0.18);
        }

        .stButton > button:disabled,
        .stFormSubmitButton > button:disabled {
            background: #CBD5E1;
            border-color: #CBD5E1;
            color: #64748B;
            box-shadow: none;
            transform: none;
        }

        div[data-testid="stTabs"] button {
            font-weight: 800;
            color: var(--as-muted);
        }

        div[data-testid="stTabs"] button[aria-selected="true"] {
            color: var(--as-blue);
        }

        div[data-testid="stDataFrame"] {
            border: 1px solid var(--as-border);
            border-radius: var(--as-radius);
            overflow: hidden;
            box-shadow: 0 4px 16px rgba(15, 76, 129, 0.04);
        }

        input, textarea, select {
            border-radius: 12px !important;
        }

        div[data-baseweb="input"],
        div[data-baseweb="select"],
        div[data-baseweb="textarea"] {
            border-radius: 12px;
        }

        .as-card,
        .pj-card {
            background: var(--as-card);
            border: 1px solid var(--as-border);
            border-radius: var(--as-radius);
            padding: 1.25rem;
            box-shadow: var(--as-shadow);
            margin-bottom: 1rem;
        }

        .as-page-header,
        .pj-page-header {
            background: linear-gradient(135deg, #FFFFFF 0%, var(--as-soft-blue) 100%);
            border: 1px solid var(--as-border);
            border-radius: 22px;
            padding: 1.4rem 1.6rem;
            margin-bottom: 1.25rem;
            box-shadow: var(--as-shadow);
        }

        .as-page-header h1,
        .pj-page-header h1 {
            margin: 0;
            color: var(--as-blue);
            font-size: 2rem;
            line-height: 1.2;
        }

        .as-page-header p,
        .pj-page-header p {
            margin: 0.35rem 0 0 0;
            color: var(--as-muted);
            font-size: 1rem;
        }

        .as-module-card,
        .pj-module-card {
            background: var(--as-card);
            border: 1px solid var(--as-border);
            border-radius: 20px;
            padding: 1.25rem;
            min-height: 150px;
            box-shadow: var(--as-shadow);
            transition: all 0.15s ease-in-out;
            margin-bottom: 0.75rem;
        }

        .as-module-card:hover,
        .pj-module-card:hover {
            transform: translateY(-2px);
            box-shadow: 0 12px 28px rgba(15, 76, 129, 0.12);
        }

        .as-module-card h3,
        .pj-module-card h3 {
            color: var(--as-blue);
            margin-top: 0;
            margin-bottom: 0.4rem;
            font-size: 1.15rem;
        }

        .as-module-card p,
        .as-info-card p,
        .as-metric-card p,
        .pj-module-card p,
        .pj-info-card p,
        .pj-metric-card p {
            color: var(--as-muted);
            margin-bottom: 0;
        }

        .as-info-card,
        .pj-info-card {
            background: var(--as-card);
            border: 1px solid var(--as-border);
            border-left: 5px solid var(--as-teal);
            border-radius: var(--as-radius);
            padding: 1rem 1.15rem;
            box-shadow: 0 4px 16px rgba(15, 76, 129, 0.05);
            margin-bottom: 1rem;
        }

        .as-info-card h3,
        .pj-info-card h3 {
            color: var(--as-blue);
            margin: 0 0 0.35rem 0;
            font-size: 1rem;
        }

        .as-metric-card,
        .pj-metric-card {
            background: var(--as-card);
            border: 1px solid var(--as-border);
            border-radius: var(--as-radius);
            padding: 1rem 1.15rem;
            box-shadow: var(--as-shadow);
            margin-bottom: 1rem;
        }

        .as-metric-card .value,
        .pj-metric-card .value {
            color: var(--as-blue);
            font-size: 1.8rem;
            font-weight: 800;
            line-height: 1.1;
        }

        .as-brand-card {
            background: linear-gradient(135deg, var(--as-blue) 0%, var(--as-blue-hover) 100%);
            color: #FFFFFF;
            border-radius: 20px;
            padding: 1rem;
            margin-bottom: 1rem;
            box-shadow: var(--as-shadow);
        }

        .as-brand-card h3 {
            color: #FFFFFF;
            margin: 0;
            font-size: 1.1rem;
        }

        .as-brand-card p {
            color: rgba(255,255,255,0.82);
            margin: 0.25rem 0 0 0;
            font-size: 0.9rem;
        }

        .as-current-page-pill {
            display: inline-block;
            background: var(--as-soft-blue);
            color: var(--as-blue);
            border: 1px solid var(--as-border);
            border-radius: 999px;
            padding: 0.35rem 0.75rem;
            font-size: 0.85rem;
            font-weight: 800;
            margin: 0.25rem 0 0.8rem 0;
        }

        .as-accent-teal {
            border-left: 5px solid var(--as-teal);
        }

        .as-muted,
        .pj-muted {
            color: var(--as-muted);
        }

        .pj-divider {
            height: 1px;
            background: var(--as-border);
            margin: 1.25rem 0;
        }
        </style>
        ''',
        unsafe_allow_html=True,
    )


st.set_page_config(page_title=APP_NAME, page_icon=get_page_icon(), layout="wide")
inject_custom_css()


def render_page_header(title, subtitle=None):
    title = html.escape(str(title))
    subtitle_html = f"<p>{html.escape(str(subtitle))}</p>" if subtitle else ""
    st.markdown(
        f'<div class="as-page-header"><h1>{title}</h1>{subtitle_html}</div>',
        unsafe_allow_html=True,
    )


def render_app_brand(compact=False):
    if os.path.exists(APP_ICON_PATH):
        st.image(APP_ICON_PATH, width=72 if compact else 96)
    if compact:
        st.markdown(
            f'<div class="as-brand-card"><h3>{html.escape(APP_NAME)}</h3><p>Asistente clinico con IA</p></div>',
            unsafe_allow_html=True,
        )


def render_info_card(title, body, accent=None):
    accent_style = f' style="border-left-color:{accent};"' if accent else ""
    title = html.escape(str(title))
    body = html.escape(str(body))
    st.markdown(
        f'<div class="as-info-card"{accent_style}><h3>{title}</h3><p>{body}</p></div>',
        unsafe_allow_html=True,
    )


def render_metric_card(label, value, helper=None):
    label = html.escape(str(label))
    value = html.escape(str(value))
    helper_html = f"<p>{html.escape(str(helper))}</p>" if helper else ""
    st.markdown(
        f'<div class="as-metric-card"><p>{label}</p><div class="value">{value}</div>{helper_html}</div>',
        unsafe_allow_html=True,
    )


def render_nav_card(title, description, button_label, page_name, key):
    title = html.escape(str(title))
    description = html.escape(str(description))
    st.markdown(
        f'<div class="as-module-card"><h3>{title}</h3><p>{description}</p></div>',
        unsafe_allow_html=True,
    )
    if st.button(button_label, key=key):
        st.session_state['active_page'] = page_name
        st.rerun()


def get_active_page_label():
    current_active = st.session_state.get('active_page') or 'Menu principal'
    if current_active in ['Pacientes', 'Ficha paciente']:
        return 'Historial atenciones de pacientes'
    return current_active


def render_sidebar_nav_button(label, page_name, key):
    active = st.session_state.get('active_page') == page_name
    if page_name == 'Pacientes':
        active = st.session_state.get('active_page') in ['Pacientes', 'Ficha paciente']
    if st.button(label, key=key, use_container_width=True, type='primary' if active else 'secondary'):
        st.session_state['active_page'] = page_name
        st.rerun()

DATA_DIR = os.path.join(BASE_DIR, "data")
SESSIONS_DIR = os.path.join(DATA_DIR, "sessions")
EXPORTS_DIR = os.path.join(DATA_DIR, "exports")
DB_PATH = os.path.join(DATA_DIR, "proyecto_jorge_ambulatorio.sqlite3")
CIE10_SEED_PATH = os.path.join(BASE_DIR, "cie10_seed.csv")

for d in [DATA_DIR, SESSIONS_DIR, EXPORTS_DIR]:
    os.makedirs(d, exist_ok=True)

SCHEMA_VERSION = "v5-ambulatory-ehr-flat-schema-hl7"
PROMPT_VERSION = "v3-jorge-ambulatory-icd10"
LLM_MODEL = get_setting("OPENAI_NOTE_MODEL", "gpt-4o")
VISION_MODEL = get_setting("OPENAI_VISION_MODEL", "gpt-4o-mini")
CHAT_MODEL = get_setting("OPENAI_CHAT_MODEL", "gpt-4o-mini")
SMTP_HOST = get_setting("SMTP_HOST", "smtp.gmail.com")
SMTP_PORT = get_int_setting("SMTP_PORT", 587)
SMTP_USER = get_setting("SMTP_USER", "")
SMTP_PASSWORD = get_setting("SMTP_PASSWORD", "")
SMTP_FROM = get_setting("SMTP_FROM", SMTP_USER)
SMTP_ENABLED = get_bool_setting("SMTP_ENABLED", False)
AUTO_EMAIL_ON_NOTE = get_bool_setting("AUTO_EMAIL_ON_NOTE", False)
HL7_SETTINGS = HL7Settings.from_env()

AUDIO_SR = 16000
MAX_CONTEXT_ORGANIZED = 18000
MAX_CONTEXT_IMAGES = 6000
DEFAULT_AUDIO_DEVICE = get_setting("WINDOWS_AUDIO_DEVICE", "")

# =========================================================
# Utilidades generales
# =========================================================
def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def hl7_ts(dt: Optional[datetime] = None) -> str:
    return (dt or datetime.now()).strftime("%Y%m%d%H%M%S")


def sha1_text(s: str) -> str:
    return hashlib.sha1(s.encode("utf-8")).hexdigest()


def save_json(path: str, data: dict):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def load_json(path: str) -> Optional[dict]:
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return None


def export_txt(content: str, filename: str) -> str:
    path = os.path.join(EXPORTS_DIR, filename)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content or "")
    return path


def export_docx(title: str, content: str, filename: str) -> str:
    path = os.path.join(EXPORTS_DIR, filename)
    doc = Document()
    doc.add_heading(title, level=1)
    for line in (content or "").splitlines():
        doc.add_paragraph(line)
    doc.save(path)
    return path


def export_fhir_bundle_json(
    patient: dict,
    encounter: dict,
    diagnoses_df: Optional[pd.DataFrame],
    reviewed_note_text: str,
    note_json: Optional[dict],
    auth_user: Optional[dict],
) -> str:
    from app.core.clinical_note import clinical_note_from_project_rows
    from app.fhir.bundles import build_document_bundle, export_bundle_json

    diagnoses = diagnoses_df.to_dict("records") if diagnoses_df is not None and not diagnoses_df.empty else []
    clinical_note = clinical_note_from_project_rows(
        patient,
        encounter,
        diagnoses=diagnoses,
        note_json=note_json,
        reviewed_note_text=reviewed_note_text,
        clinician_specialty=(auth_user or {}).get("specialty"),
        clinician_license_number=(auth_user or {}).get("professional_id"),
        physician_validated=(encounter.get("status") == "closed"),
    )
    bundle = build_document_bundle(clinical_note)
    filename = f"fhir_bundle_{clinical_note.note_id}.json"
    return export_bundle_json(bundle, os.path.join(EXPORTS_DIR, filename))


def clean_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()

def normalize_facility_code(value: str) -> str:
    code = re.sub(r"[^A-Za-z0-9]", "", value or "").upper()
    if not code:
        raise ValueError("La sigla del centro es obligatoria.")
    if len(code) > 10:
        raise ValueError("La sigla del centro debe tener 10 caracteres o menos.")
    return code

def suggest_facility_code(name: str) -> str:
    words = re.findall(r"[A-Za-z0-9]+", name or "")
    if not words:
        return "CTR"
    letters = "".join(w[0] for w in words if not w.isdigit())
    numbers = "".join(w for w in words if w.isdigit())
    return normalize_facility_code((letters or words[0][:3]) + numbers)

# =========================================================
# Base de datos SQLite
# =========================================================
def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS patients (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                mrn TEXT UNIQUE,
                first_name TEXT,
                last_name TEXT,
                national_id TEXT,
                birth_date TEXT,
                sex TEXT,
                phone TEXT,
                email TEXT,
                address TEXT,
                guardian_name TEXT,
                facility_id INTEGER,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS encounters (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                patient_id INTEGER NOT NULL,
                visit_type TEXT,
                brought_by TEXT,
                informant TEXT,
                provider_name TEXT,
                facility_name TEXT,
                chief_complaint TEXT,
                note_text TEXT,
                note_json TEXT,
                transcript_text TEXT,
                dictation_text TEXT,
                images_text TEXT,
                status TEXT DEFAULT 'draft',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(patient_id) REFERENCES patients(id)
            );

            CREATE TABLE IF NOT EXISTS diagnoses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                encounter_id INTEGER NOT NULL,
                code TEXT NOT NULL,
                description TEXT NOT NULL,
                diagnosis_type TEXT DEFAULT 'principal',
                certainty TEXT DEFAULT 'sospecha',
                created_at TEXT NOT NULL,
                FOREIGN KEY(encounter_id) REFERENCES encounters(id)
            );

            CREATE TABLE IF NOT EXISTS cie10_catalog (
                code TEXT PRIMARY KEY,
                description TEXT NOT NULL,
                chapter TEXT
            );

            CREATE TABLE IF NOT EXISTS hl7_exports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                encounter_id INTEGER NOT NULL,
                message_type TEXT NOT NULL,
                file_path TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(encounter_id) REFERENCES encounters(id)
            );
            
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'doctor', full_name TEXT, phone TEXT, specialty TEXT, professional_id TEXT,
                preferred_note_format TEXT DEFAULT 'SOAP', first_login_completed INTEGER DEFAULT 0, active INTEGER DEFAULT 1,
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS facilities (
                id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE NOT NULL, code TEXT UNIQUE, address TEXT, phone TEXT, active INTEGER DEFAULT 1,
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS appointments (
                id INTEGER PRIMARY KEY AUTOINCREMENT, patient_id INTEGER, facility_id INTEGER, provider_user_id INTEGER,
                appointment_date TEXT NOT NULL, appointment_time TEXT NOT NULL, duration_minutes INTEGER DEFAULT 30,
                status TEXT DEFAULT 'agendada', reason TEXT, encounter_id INTEGER, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            """
        )
        ensure_hl7_tables(conn)
        patient_cols = [r["name"] for r in conn.execute("PRAGMA table_info(patients)").fetchall()]
        if "facility_id" not in patient_cols:
            conn.execute("ALTER TABLE patients ADD COLUMN facility_id INTEGER")
        facility_cols = [r["name"] for r in conn.execute("PRAGMA table_info(facilities)").fetchall()]
        if "code" not in facility_cols:
            conn.execute("ALTER TABLE facilities ADD COLUMN code TEXT")
        conn.commit()


def seed_cie10_if_empty():
    with db() as conn:
        count = conn.execute("SELECT COUNT(*) AS n FROM cie10_catalog").fetchone()["n"]
        if count > 0:
            return
        rows = []
        if os.path.exists(CIE10_SEED_PATH):
            df = pd.read_csv(CIE10_SEED_PATH)
            for _, r in df.iterrows():
                rows.append((clean_str(r.get("code")), clean_str(r.get("description")), clean_str(r.get("chapter"))))
        if not rows:
            rows = [
                ("G40", "Epilepsia", "VI Enfermedades del sistema nervioso"),
                ("G40.9", "Epilepsia, no especificada", "VI Enfermedades del sistema nervioso"),
                ("R56.8", "Otras convulsiones y las no especificadas", "XVIII Síntomas y signos"),
                ("F80.9", "Trastorno del desarrollo del habla y del lenguaje, no especificado", "V Trastornos mentales y del comportamiento"),
                ("F84.0", "Autismo infantil", "V Trastornos mentales y del comportamiento"),
                ("G43", "Migraña", "VI Enfermedades del sistema nervioso"),
                ("R51", "Cefalea", "XVIII Síntomas y signos"),
                ("G80", "Parálisis cerebral", "VI Enfermedades del sistema nervioso"),
                ("F90.0", "Perturbación de la actividad y de la atención", "V Trastornos mentales y del comportamiento"),
                ("Z00.1", "Control de salud de rutina del niño", "XXI Factores que influyen en el estado de salud"),
            ]
        conn.executemany(
            "INSERT OR IGNORE INTO cie10_catalog(code, description, chapter) VALUES (?, ?, ?)",
            rows,
        )
        conn.commit()


def create_or_update_patient(data: dict) -> int:
    facility_id = data.get("facility_id")
    if not facility_id:
        raise ValueError("Debes seleccionar el centro al que pertenece el paciente.")
    ts = now_iso()
    with db() as conn:
        patient_id = data.get("id")
        mrn = data.get("mrn")
        existing = conn.execute("SELECT id, mrn FROM patients WHERE id=?", (patient_id,)).fetchone() if patient_id else None
        if not existing and mrn:
            existing = conn.execute("SELECT id, mrn FROM patients WHERE mrn=?", (mrn,)).fetchone()
        if existing:
            pid = existing["id"]
            conn.execute(
                """
                UPDATE patients SET first_name=?, last_name=?, national_id=?, birth_date=?, sex=?, phone=?, email=?,
                address=?, guardian_name=?, facility_id=?, updated_at=? WHERE id=?
                """,
                (
                    data.get("first_name"), data.get("last_name"), data.get("national_id"), data.get("birth_date"),
                    data.get("sex"), data.get("phone"), data.get("email"), data.get("address"), data.get("guardian_name"), facility_id, ts, pid,
                ),
            )
        else:
            mrn = generate_patient_mrn(conn, int(facility_id))
            cur = conn.execute(
                """
                INSERT INTO patients(mrn, first_name, last_name, national_id, birth_date, sex, phone, email, address,
                guardian_name, facility_id, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    mrn, data.get("first_name"), data.get("last_name"), data.get("national_id"), data.get("birth_date"),
                    data.get("sex"), data.get("phone"), data.get("email"), data.get("address"), data.get("guardian_name"), facility_id, ts, ts,
                ),
            )
            pid = cur.lastrowid
        conn.commit()
        return pid

def generate_patient_mrn(conn: sqlite3.Connection, facility_id: int) -> str:
    facility = conn.execute("SELECT code, name FROM facilities WHERE id=?", (facility_id,)).fetchone()
    if not facility:
        raise ValueError("Centro no encontrado para generar el identificador del paciente.")
    code = normalize_facility_code(facility["code"] or suggest_facility_code(facility["name"]))
    rows = conn.execute("SELECT mrn FROM patients WHERE facility_id=? AND mrn LIKE ?", (facility_id, f"{code}%")).fetchall()
    max_seq = 0
    for row in rows:
        suffix = str(row["mrn"] or "")[len(code):]
        if suffix.isdigit():
            max_seq = max(max_seq, int(suffix))
    return f"{code}{max_seq + 1:05d}"

def preview_next_patient_mrn(facility_id: int) -> str:
    with db() as conn:
        return generate_patient_mrn(conn, facility_id)


def create_encounter(patient_id: int, data: dict) -> int:
    ts = now_iso()
    with db() as conn:
        cur = conn.execute(
            """
            INSERT INTO encounters(patient_id, visit_type, brought_by, informant, provider_name, facility_name,
            chief_complaint, status, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, 'draft', ?, ?)
            """,
            (
                patient_id, data.get("visit_type"), data.get("brought_by"), data.get("informant"),
                data.get("provider_name"), data.get("facility_name"), data.get("chief_complaint"), ts, ts,
            ),
        )
        conn.commit()
        return cur.lastrowid


def update_encounter_content(encounter_id: int, **kwargs):
    allowed = ["note_text", "note_json", "transcript_text", "dictation_text", "images_text", "chief_complaint", "status"]
    pairs = []
    values = []
    for k, v in kwargs.items():
        if k in allowed:
            pairs.append(f"{k}=?")
            values.append(v)
    if not pairs:
        return
    pairs.append("updated_at=?")
    values.append(now_iso())
    values.append(encounter_id)
    with db() as conn:
        conn.execute(f"UPDATE encounters SET {', '.join(pairs)} WHERE id=?", values)
        conn.commit()


def add_diagnosis(encounter_id: int, code: str, description: str, diagnosis_type: str, certainty: str):
    with db() as conn:
        conn.execute(
            "INSERT INTO diagnoses(encounter_id, code, description, diagnosis_type, certainty, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            (encounter_id, code.strip().upper(), description.strip(), diagnosis_type, certainty, now_iso()),
        )
        conn.commit()


def delete_diagnosis(diag_id: int):
    with db() as conn:
        conn.execute("DELETE FROM diagnoses WHERE id=?", (diag_id,))
        conn.commit()


def get_patients_df() -> pd.DataFrame:
    with db() as conn:
        return pd.read_sql_query(
            """
            SELECT p.id, p.mrn, p.first_name, p.last_name, p.national_id, p.birth_date, p.sex,
                   p.guardian_name, p.facility_id, COALESCE(f.name, 'Sin centro') AS centro, p.updated_at
            FROM patients p
            LEFT JOIN facilities f ON f.id = p.facility_id
            ORDER BY centro, p.updated_at DESC
            """,
            conn,
        )


def get_patient(patient_id: int) -> Optional[dict]:
    with db() as conn:
        row = conn.execute(
            "SELECT p.*, COALESCE(f.name, '') AS facility_name FROM patients p LEFT JOIN facilities f ON f.id=p.facility_id WHERE p.id=?",
            (patient_id,),
        ).fetchone()
        return dict(row) if row else None


def get_patient_by_mrn(mrn: str) -> Optional[dict]:
    with db() as conn:
        row = conn.execute("SELECT * FROM patients WHERE mrn=?", (mrn.strip(),)).fetchone()
        return dict(row) if row else None


def get_patient_by_national_id(national_id: str) -> Optional[dict]:
    with db() as conn:
        row = conn.execute("SELECT * FROM patients WHERE national_id=?", (national_id.strip(),)).fetchone()
        return dict(row) if row else None


def search_patients(term: str, limit: int = 20) -> pd.DataFrame:
    q = f"%{(term or '').strip()}%"
    with db() as conn:
        return pd.read_sql_query(
            """
            SELECT p.id, p.mrn, p.first_name, p.last_name, p.national_id, p.birth_date, p.sex,
                   p.phone, p.email, p.address, p.guardian_name, p.facility_id, COALESCE(f.name, 'Sin centro') AS centro, p.updated_at
            FROM patients p
            LEFT JOIN facilities f ON f.id = p.facility_id
            WHERE p.mrn LIKE ? OR p.national_id LIKE ? OR p.first_name LIKE ? OR p.last_name LIKE ?
               OR p.birth_date LIKE ? OR p.sex LIKE ? OR p.phone LIKE ? OR p.email LIKE ?
               OR p.address LIKE ? OR p.guardian_name LIKE ? OR f.name LIKE ?
            ORDER BY centro, p.updated_at DESC LIMIT ?
            """,
            conn,
            params=(q, q, q, q, q, q, q, q, q, q, q, limit),
        )


def get_encounters_df(patient_id: Optional[int] = None) -> pd.DataFrame:
    with db() as conn:
        if patient_id:
            return pd.read_sql_query(
                "SELECT id, patient_id, visit_type, chief_complaint, status, created_at, updated_at FROM encounters WHERE patient_id=? ORDER BY created_at DESC",
                conn,
                params=(patient_id,),
            )
        return pd.read_sql_query(
            "SELECT id, patient_id, visit_type, chief_complaint, status, created_at, updated_at FROM encounters ORDER BY created_at DESC",
            conn,
        )

def get_patient_notes_df(patient_id: int) -> pd.DataFrame:
    with db() as conn:
        return pd.read_sql_query(
            """
            SELECT id, visit_type, chief_complaint, status, created_at, updated_at, note_text, note_json
            FROM encounters
            WHERE patient_id=?
            ORDER BY created_at DESC
            """,
            conn,
            params=(patient_id,),
        )


def get_encounter(encounter_id: int) -> Optional[dict]:
    with db() as conn:
        row = conn.execute("SELECT * FROM encounters WHERE id=?", (encounter_id,)).fetchone()
        return dict(row) if row else None


def get_diagnoses_df(encounter_id: int) -> pd.DataFrame:
    with db() as conn:
        return pd.read_sql_query(
            "SELECT id, code, description, diagnosis_type, certainty, created_at FROM diagnoses WHERE encounter_id=? ORDER BY id",
            conn,
            params=(encounter_id,),
        )


def search_cie10(term: str, limit: int = 30) -> pd.DataFrame:
    q = f"%{(term or '').strip()}%"
    with db() as conn:
        return pd.read_sql_query(
            """
            SELECT code, description, chapter FROM cie10_catalog
            WHERE code LIKE ? OR description LIKE ?
            ORDER BY code LIMIT ?
            """,
            conn,
            params=(q, q, limit),
        )

# =========================================================
# Audio y cache
# =========================================================
AUDIO_UPLOAD_TYPES = ["wav", "mp3", "m4a", "ogg"]


def patient_audio_prefix(patient_id: Optional[int]) -> str:
    return f"{int(patient_id)}_paciente" if patient_id else ""


def next_patient_audio_path(patient_id: Optional[int], kind: str, extension: str = "wav") -> str:
    if not patient_id:
        raise ValueError("Primero selecciona o guarda un paciente para asociar el audio.")
    prefix = patient_audio_prefix(patient_id)
    existing = glob.glob(os.path.join(SESSIONS_DIR, f"{prefix}_audio*.*"))
    max_seq = 0
    pattern = re.compile(rf"^{re.escape(prefix)}_audio(\d+)", re.IGNORECASE)
    for path in existing:
        match = pattern.match(os.path.basename(path))
        if match:
            max_seq = max(max_seq, int(match.group(1)))
    safe_kind = re.sub(r"[^A-Za-z0-9_-]+", "", kind or "audio").lower()
    safe_ext = re.sub(r"[^A-Za-z0-9]+", "", extension or "wav").lower()
    if safe_ext not in AUDIO_UPLOAD_TYPES:
        safe_ext = "wav"
    return os.path.join(SESSIONS_DIR, f"{prefix}_audio{max_seq + 1:02d}_{safe_kind}.{safe_ext}")


def patient_audio_path(patient_id: Optional[int], kind: str) -> str:
    return next_patient_audio_path(patient_id, kind, "wav")


def save_uploaded_audio(uploaded_file, patient_id: Optional[int], kind: str) -> str:
    original_name = uploaded_file.name or ""
    extension = os.path.splitext(original_name)[1].lstrip(".").lower() or "wav"
    data = uploaded_file.getvalue()
    upload_key = sha1_text(f"{patient_id}|{kind}|{original_name}|{len(data)}|{hashlib.sha1(data).hexdigest()}")
    cached = st.session_state.get("uploaded_audio_cache") or {}
    if cached.get("key") == upload_key and cached.get("path") and os.path.exists(cached["path"]):
        return cached["path"]
    path = next_patient_audio_path(patient_id, f"{kind}_subido", extension)
    with open(path, "wb") as f:
        f.write(data)
    st.session_state["uploaded_audio_cache"] = {"key": upload_key, "path": path}
    return path


def audio_belongs_to_patient(path: str, patient_id: Optional[int]) -> bool:
    if not patient_id:
        return False
    return os.path.basename(path).lower().startswith(f"{patient_audio_prefix(patient_id).lower()}_audio")


def get_latest_audio(prefix: Optional[str] = None, patient_id: Optional[int] = None) -> Optional[str]:
    patterns = [f"*.{ext}" if not prefix else f"{prefix}_*.{ext}" for ext in AUDIO_UPLOAD_TYPES]
    files = []
    for pattern in patterns:
        files.extend(glob.glob(os.path.join(SESSIONS_DIR, pattern)))
    if patient_id:
        files = [p for p in files if audio_belongs_to_patient(p, patient_id)]
    if not files:
        return None
    files.sort(key=os.path.getmtime, reverse=True)
    return files[0]


def cache_key_for_file(path: str) -> str:
    st_info = os.stat(path)
    key = f"{path}|{st_info.st_size}|{int(st_info.st_mtime)}"
    return sha1_text(key)


def transcript_cache_path(key: str, kind: str) -> str:
    return os.path.join(SESSIONS_DIR, f"transcript_{kind}_{key}.json")


def audio_file_label(path: Optional[str]) -> str:
    if not path:
        return ''
    return os.path.basename(path)


def list_audio_files(patient_id: Optional[int] = None) -> Tuple[List[str], List[str], List[str]]:
    all_wavs = []
    for ext in AUDIO_UPLOAD_TYPES:
        all_wavs.extend(glob.glob(os.path.join(SESSIONS_DIR, f"*.{ext}")))
    if patient_id:
        all_wavs = [p for p in all_wavs if audio_belongs_to_patient(p, patient_id)]
    all_wavs.sort(key=os.path.getmtime, reverse=True)
    conv_files = [p for p in all_wavs if 'conversacion' in os.path.basename(p).lower() or os.path.basename(p).lower().startswith('conv_')]
    dict_files = [p for p in all_wavs if 'dictado' in os.path.basename(p).lower() or os.path.basename(p).lower().startswith('dict_')]
    other_files = [p for p in all_wavs if p not in conv_files and p not in dict_files]
    return conv_files, dict_files, other_files


def note_cache_path(note_key: str) -> str:
    return os.path.join(SESSIONS_DIR, f"note_{note_key}.json")


def list_audio_devices() -> str:
    if not SOUNDDEVICE_AVAILABLE:
        return "La grabacion local no esta disponible porque sounddevice/PortAudio no se pudo cargar en este entorno."
    try:
        devices = sd.query_devices()
        lines = [f"Sistema: {platform.system()} {platform.release()}", "", "Dispositivos de audio detectados:"]
        for idx, dev in enumerate(devices):
            max_in = int(dev.get("max_input_channels", 0) or 0)
            max_out = int(dev.get("max_output_channels", 0) or 0)
            default_sr = dev.get("default_samplerate", "")
            marker = " <-- entrada" if max_in > 0 else ""
            lines.append(f"[{idx}] {dev.get('name')} | in={max_in} out={max_out} sr={default_sr}{marker}")
        lines.append("")
        lines.append("Si necesitas forzar un microfono, configura WINDOWS_AUDIO_DEVICE con el numero del dispositivo, por ejemplo WINDOWS_AUDIO_DEVICE=1")
        return "\n".join(lines)
    except Exception as e:
        return f"No se pudieron listar dispositivos de audio: {e}"


def _audio_device_value():
    raw = (DEFAULT_AUDIO_DEVICE or "").strip()
    if raw == "":
        return None
    try:
        return int(raw)
    except ValueError:
        return raw


def record_wav(path: str, seconds: int, samplerate: int = AUDIO_SR, audio_device=None):
    if not SOUNDDEVICE_AVAILABLE:
        raise RuntimeError(
            "La grabacion local no esta disponible en este entorno porque sounddevice/PortAudio no se pudo cargar. "
            "En Streamlit Cloud usa la opcion para subir un audio."
        )
    device = audio_device if audio_device is not None else _audio_device_value()
    try:
        audio = sd.rec(
            int(seconds * samplerate),
            samplerate=samplerate,
            channels=1,
            dtype="float32",
            device=device,
        )
        sd.wait()
        sf.write(path, audio, samplerate)
        return path
    except Exception as e:
        raise RuntimeError(
            "No se pudo grabar audio en Windows.\n\n"
            f"Detalle: {e}\n\n"
            "Revisa que el micrófono esté permitido para aplicaciones de escritorio en Windows, "
            "que no esté ocupado por otra aplicación y, si es necesario, usa 'Listar dispositivos de audio' "
            "para configurar la variable WINDOWS_AUDIO_DEVICE."
        )


def transcribe_local_with_cache(wav_path: str, kind: str) -> Tuple[str, List[Dict]]:
    key = cache_key_for_file(wav_path)
    cpath = transcript_cache_path(key, kind)
    cached = load_json(cpath)
    if cached:
        return cached.get("text", ""), cached.get("segments", [])
    text, segments = transcribe_whisper(wav_path)
    save_json(cpath, {"text": text, "segments": segments})
    return text, segments

# =========================================================
# Schema IA plano sin $ref
# =========================================================
def field_value_schema() -> dict:
    return {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "value": {"type": ["string", "null"]},
            "evidence": {
                "type": ["object", "null"],
                "additionalProperties": False,
                "properties": {
                    "quote": {"type": "string"},
                    "start_sec": {"type": ["number", "null"]},
                    "end_sec": {"type": ["number", "null"]},
                },
                "required": ["quote", "start_sec", "end_sec"],
            },
        },
        "required": ["value", "evidence"],
    }


def clinical_note_flat_schema() -> dict:
    fv = field_value_schema()
    props = {
        "visit_type": {"type": "string"},
        "patient_name": fv,
        "patient_age": fv,
        "patient_sex": fv,
        "brought_by": fv,
        "informant": fv,
        "chief_complaint": fv,
        "hpi_onset": fv,
        "hpi_course": fv,
        "hpi_frequency": fv,
        "hpi_duration": fv,
        "hpi_triggers": fv,
        "hpi_relief": fv,
        "neuro_symptoms": {"type": "array", "items": {"type": "string"}},
        "past_medical_history": fv,
        "allergies": fv,
        "meds": fv,
        "family": {
            "type": "object", "additionalProperties": False,
            "properties": {"mother_history": fv, "father_history": fv, "consanguinity": fv, "siblings": fv},
            "required": ["mother_history", "father_history", "consanguinity", "siblings"],
        },
        "perinatal": {
            "type": "object", "additionalProperties": False,
            "properties": {
                "pregnancy_controlled": fv, "pregnancy_control_from_week": fv, "pregnancy_diseases": fv,
                "fetal_movements": fv, "delivery_type": fv, "delivery_complications": fv,
                "fetal_distress": fv, "gestational_age_weeks": fv, "birth_weight": fv,
                "birth_length": fv, "birth_head_circumference": fv, "neonatal_disease": fv,
                "discharged_with_mother": fv, "neonatal_hospitalization_reason": fv,
            },
            "required": [
                "pregnancy_controlled", "pregnancy_control_from_week", "pregnancy_diseases", "fetal_movements",
                "delivery_type", "delivery_complications", "fetal_distress", "gestational_age_weeks", "birth_weight",
                "birth_length", "birth_head_circumference", "neonatal_disease", "discharged_with_mother",
                "neonatal_hospitalization_reason",
            ],
        },
        "development": {
            "type": "object", "additionalProperties": False,
            "properties": {
                "delay_gross_motor": fv, "delay_fine_motor": fv, "delay_social": fv,
                "delay_language": fv, "milestones": fv,
            },
            "required": ["delay_gross_motor", "delay_fine_motor", "delay_social", "delay_language", "milestones"],
        },
        "exam_neuro": fv,
        "exams_reviewed": fv,
        "diagnostic_hypothesis": fv,
        "suspicion_or_confirmed": fv,
        "suggested_icd10": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "code": {"type": "string"},
                    "description": {"type": "string"},
                    "certainty": {"type": "string"},
                    "reason": {"type": "string"},
                },
                "required": ["code", "description", "certainty", "reason"],
            },
        },
        "indications": {
            "type": "object", "additionalProperties": False,
            "properties": {"general": fv, "school": fv, "meds": fv, "followup": fv, "referrals": fv},
            "required": ["general", "school", "meds", "followup", "referrals"],
        },
        "red_flags_mentioned": {"type": "array", "items": {"type": "string"}},
        "missing_questions": {"type": "array", "items": {"type": "string"}},
        "contradictions": {"type": "array", "items": {"type": "string"}},
    }
    return {"type": "object", "additionalProperties": False, "properties": props, "required": list(props.keys())}


def extract_structured(payload: dict) -> dict:
    schema = clinical_note_flat_schema()
    system = SYSTEM_INSTRUCTIONS + """

EXTENSION AMBULATORIA:
- Además de estructurar la nota, sugiere diagnósticos CIE-10 SOLO si hay evidencia clínica explícita.
- Los diagnósticos sugeridos no reemplazan el juicio médico. Si no corresponde, devuelve suggested_icd10 como array vacío.
- Mantén baja alucinación: no inventes códigos exactos si no estás razonablemente seguro; usa descripciones conservadoras.
"""
    resp = get_openai_client().responses.create(
        model=LLM_MODEL,
        temperature=0.1,
        input=[
            {"role": "system", "content": system},
            {"role": "user", "content": "DATA:\n" + json.dumps(payload, ensure_ascii=False)},
        ],
        text={"format": {"type": "json_schema", "name": "clinical_note", "schema": schema, "strict": True}},
    )
    return json.loads(resp.output_text)

# =========================================================
# Render nota
# =========================================================
def safe_get(note: dict, path: str) -> Optional[dict]:
    cur = note
    for p in path.split("."):
        if not isinstance(cur, dict) or p not in cur:
            return None
        cur = cur[p]
    return cur if isinstance(cur, dict) else None


def fv(note: dict, path: str) -> Optional[str]:
    obj = safe_get(note, path)
    return obj.get("value") if obj else None


def render_note(note: dict, meta: dict, diagnoses_df: Optional[pd.DataFrame] = None) -> str:
    lines: List[str] = []
    lines.append("FICHA CLINICA AMBULATORIA - NEUROLOGIA PEDIATRICA")
    lines.append("")

    header = []
    for label, value in [
        ("Tipo", note.get("visit_type") or meta.get("visit_type")),
        ("Paciente", fv(note, "patient_name") or meta.get("patient_name")),
        ("Edad", fv(note, "patient_age") or meta.get("patient_age")),
        ("Sexo", fv(note, "patient_sex") or meta.get("patient_sex")),
        ("Traído por", fv(note, "brought_by") or meta.get("brought_by")),
        ("Informante", fv(note, "informant") or meta.get("informant")),
    ]:
        if value:
            header.append(f"{label}: {value}")
    lines.append(" | ".join(header) if header else "Datos de encabezado: No mencionados")
    lines.append("")

    sections = [
        ("Motivo de consulta", fv(note, "chief_complaint")),
    ]
    for title, content in sections:
        lines.append(f"{title}:")
        lines.append(content or "No mencionado")
        lines.append("")

    lines.append("Anamnesis próxima:")
    any_hpi = False
    for label, key in [
        ("Inicio", "hpi_onset"), ("Curso", "hpi_course"), ("Frecuencia", "hpi_frequency"),
        ("Duración", "hpi_duration"), ("Gatillantes", "hpi_triggers"), ("Alivio", "hpi_relief"),
    ]:
        v = fv(note, key)
        if v:
            any_hpi = True
            lines.append(f"- {label}: {v}")
    if not any_hpi:
        lines.append("No mencionado")

    ns = note.get("neuro_symptoms") or []
    if ns:
        lines.append("- Síntomas neurológicos mencionados: " + ", ".join(ns))
    lines.append("")

    lines.append("Anamnesis remota:")
    any_remote = False
    for label, key in [("Antecedentes personales", "past_medical_history"), ("Alergias", "allergies"), ("Medicamentos", "meds")]:
        v = fv(note, key)
        if v:
            any_remote = True
            lines.append(f"- {label}: {v}")
    if not any_remote:
        lines.append("No mencionado")
    lines.append("")

    lines.append("Antecedentes familiares:")
    fam_any = False
    for label, key in [("Madre", "family.mother_history"), ("Padre", "family.father_history"), ("Consanguinidad", "family.consanguinity"), ("Hermanos", "family.siblings")]:
        v = fv(note, key)
        if v:
            fam_any = True
            lines.append(f"- {label}: {v}")
    if not fam_any:
        lines.append("No mencionado")
    lines.append("")

    lines.append("Antecedentes perinatales:")
    per_any = False
    for label, key in [
        ("Embarazo controlado", "perinatal.pregnancy_controlled"), ("Control desde semana", "perinatal.pregnancy_control_from_week"),
        ("Enfermedades embarazo", "perinatal.pregnancy_diseases"), ("Movimientos fetales", "perinatal.fetal_movements"),
        ("Tipo de parto", "perinatal.delivery_type"), ("Complicaciones parto", "perinatal.delivery_complications"),
        ("Sufrimiento fetal", "perinatal.fetal_distress"), ("Semanas gestación", "perinatal.gestational_age_weeks"),
        ("Peso", "perinatal.birth_weight"), ("Talla", "perinatal.birth_length"),
        ("Perímetro cefálico", "perinatal.birth_head_circumference"), ("Enfermedad RN", "perinatal.neonatal_disease"),
        ("Alta con madre", "perinatal.discharged_with_mother"), ("Hospitalización RN motivo", "perinatal.neonatal_hospitalization_reason"),
    ]:
        v = fv(note, key)
        if v:
            per_any = True
            lines.append(f"- {label}: {v}")
    if not per_any:
        lines.append("No mencionado")
    lines.append("")

    lines.append("Desarrollo psicomotor:")
    dev_any = False
    for label, key in [("Retraso motor grueso", "development.delay_gross_motor"), ("Retraso motor fino", "development.delay_fine_motor"), ("Retraso social", "development.delay_social"), ("Retraso lenguaje", "development.delay_language"), ("Hitos", "development.milestones")]:
        v = fv(note, key)
        if v:
            dev_any = True
            lines.append(f"- {label}: {v}")
    if not dev_any:
        lines.append("No mencionado")
    lines.append("")

    for title, key in [("Examen neurológico", "exam_neuro"), ("Exámenes revisados", "exams_reviewed")]:
        v = fv(note, key)
        lines.append(f"{title}:")
        lines.append(v or "No mencionado")
        lines.append("")

    lines.append("Hipótesis diagnóstica:")
    dx = fv(note, "diagnostic_hypothesis")
    soc = fv(note, "suspicion_or_confirmed")
    lines.append(dx or "No mencionado")
    if soc:
        lines.append(f"Estado: {soc}")
    lines.append("")

    lines.append("Diagnósticos CIE-10 registrados:")
    if diagnoses_df is not None and not diagnoses_df.empty:
        for _, r in diagnoses_df.iterrows():
            lines.append(f"- {r['code']} - {r['description']} ({r['diagnosis_type']}, {r['certainty']})")
    else:
        suggested = note.get("suggested_icd10") or []
        if suggested:
            lines.append("Sugeridos por IA, pendientes de validación médica:")
            for d in suggested:
                lines.append(f"- {d.get('code', '')} - {d.get('description', '')} ({d.get('certainty', '')})")
        else:
            lines.append("No registrados")
    lines.append("")

    lines.append("Indicaciones:")
    ind_any = False
    for label, key in [("Generales", "indications.general"), ("Colegio", "indications.school"), ("Medicamentos", "indications.meds"), ("Control", "indications.followup"), ("Derivaciones", "indications.referrals")]:
        v = fv(note, key)
        if v:
            ind_any = True
            lines.append(f"- {label}: {v}")
    if not ind_any:
        lines.append("No mencionado")

    rf = note.get("red_flags_mentioned") or []
    if rf:
        lines.append("")
        lines.append("Banderas rojas mencionadas:")
        lines.append(", ".join(rf))

    return "\n".join(lines)


def segments_to_organized_text(segments: List[Dict], title: str) -> str:
    lines = [title, ""]
    for i, s in enumerate(segments, start=1):
        start = float(s.get("start", 0.0) or 0.0)
        end = float(s.get("end", 0.0) or 0.0)
        text = (s.get("text") or "").strip()
        if text:
            lines.append(f"[{i:03d}] {start:.1f}s - {end:.1f}s: {text}")
    return "\n".join(lines).strip() + "\n"


def local_search_segments(segments: List[Dict], query: str, max_hits: int = 25) -> List[Dict]:
    q = (query or "").strip().lower()
    if not q:
        return []
    hits = []
    for s in segments:
        t = (s.get("text") or "").lower()
        if q in t:
            hits.append(s)
            if len(hits) >= max_hits:
                break
    return hits

# =========================================================
# Vision
# =========================================================
def summarize_images_with_llm(images_data_urls: List[str], instruction: str) -> str:
    content = [{"type": "text", "text": instruction}]
    for url in images_data_urls:
        content.append({"type": "input_image", "image_url": url})
    resp = get_openai_client().responses.create(model=VISION_MODEL, input=[{"role": "user", "content": content}], temperature=0.2)
    return (resp.output_text or "").strip()

# =========================================================
# HL7 v2 wrappers de compatibilidad
# =========================================================
def is_physician_reviewed(encounter: Optional[dict]) -> bool:
    return bool((encounter or {}).get("physician_reviewed")) or (encounter or {}).get("status") in {"reviewed", "closed", "completed"}


def build_adt_a04(patient: dict, encounter: dict, sending_app="ASISTENTE_SIMON", receiving_app="HIS_DESTINO") -> str:
    settings = HL7Settings.from_env()
    if sending_app:
        settings = HL7Settings(**{**settings.__dict__, "sending_app": sending_app})
    if receiving_app:
        settings = HL7Settings(**{**settings.__dict__, "receiving_app": receiving_app})
    return build_adt_a04_message(patient, encounter, settings=settings)


def build_mdm_t02(patient: dict, encounter: dict, diagnoses_df: pd.DataFrame, sending_app="ASISTENTE_SIMON", receiving_app="HIS_DESTINO") -> str:
    settings = HL7Settings.from_env()
    if sending_app:
        settings = HL7Settings(**{**settings.__dict__, "sending_app": sending_app})
    if receiving_app:
        settings = HL7Settings(**{**settings.__dict__, "receiving_app": receiving_app})
    payload_encounter = {**encounter, "physician_reviewed": is_physician_reviewed(encounter)}
    return build_mdm_t02_message(patient, payload_encounter, diagnoses_df, settings=settings, note_text=encounter.get("note_text") or "")


def save_hl7_message(encounter_id: int, message_type: str, message: str) -> str:
    transport = FileHL7Transport(os.path.join(EXPORTS_DIR, "hl7"))
    control_id = get_control_id(message) or f"AS{encounter_id}{int(time.time())}"
    path = transport.send(message, control_id, message_type)
    with db() as conn:
        conn.execute(
            "INSERT INTO hl7_exports(encounter_id, message_type, file_path, created_at) VALUES (?, ?, ?, ?)",
            (encounter_id, message_type, path, now_iso()),
        )
        conn.commit()
    return path


def latest_hl7_queue_message(encounter_id: int) -> Optional[dict]:
    with db() as conn:
        return get_latest_message(conn, encounter_id)


def create_hl7_draft(
    patient: dict,
    encounter: dict,
    diagnoses_df: Optional[pd.DataFrame],
    message_type: str,
    user_id: Optional[int],
) -> int:
    settings = HL7Settings.from_env()
    encounter_payload = {**encounter, "physician_reviewed": is_physician_reviewed(encounter)}
    if message_type == "ADT_A04":
        payload = build_adt_a04_message(patient, encounter_payload, settings=settings)
    else:
        payload = build_mdm_t02_message(patient, encounter_payload, diagnoses_df, settings=settings, note_text=encounter.get("note_text") or "")
        message_type = "MDM_T02"
    control_id = get_control_id(payload)
    with db() as conn:
        message_id = enqueue_message(conn, int(encounter["id"]), message_type, control_id, payload, created_by_user_id=user_id, status="DRAFT")
        audit_event(
            conn,
            action="hl7_generated",
            entity_type="hl7_message",
            entity_id=str(message_id),
            user_id=user_id,
            patient_id=patient.get("id"),
            encounter_id=encounter.get("id"),
            metadata={"message_type": message_type, "control_id": control_id, "payload_sha256": compute_payload_sha256(payload)},
        )
        conn.commit()
    return message_id


def validate_hl7_queue_message(message_id: int, encounter: dict) -> Tuple[bool, str]:
    settings = HL7Settings.from_env()
    with db() as conn:
        message = get_message(conn, message_id)
        if not message:
            return False, "Mensaje HL7 no encontrado."
        try:
            validate_message_ready(
                message["payload"],
                message["message_type"],
                settings=settings,
                note_text=encounter.get("note_text") or "",
                physician_reviewed=is_physician_reviewed(encounter),
            )
            update_message_status(conn, message_id, "READY")
            audit_event(conn, "hl7_validated", "hl7_message", str(message_id), encounter_id=encounter.get("id"), metadata={"control_id": message["control_id"]})
            conn.commit()
            return True, "Mensaje HL7 validado y marcado READY."
        except Exception as exc:
            update_message_status(conn, message_id, "ERROR", error_message=str(exc))
            audit_event(conn, "hl7_validation_failed", "hl7_message", str(message_id), encounter_id=encounter.get("id"), metadata={"error": str(exc)})
            conn.commit()
            return False, str(exc)


def export_hl7_queue_message(message_id: int, user_id: Optional[int]) -> Tuple[bool, str]:
    with db() as conn:
        message = get_message(conn, message_id)
        if not message:
            return False, "Mensaje HL7 no encontrado."
        path = save_hl7_message(int(message["encounter_id"]), message["message_type"], message["payload"])
        audit_event(conn, "hl7_downloaded", "hl7_message", str(message_id), user_id=user_id, encounter_id=message.get("encounter_id"), metadata={"control_id": message["control_id"], "path": path})
        conn.commit()
    return True, path


def send_hl7_queue_message(message_id: int, user_id: Optional[int]) -> Tuple[bool, str]:
    settings = HL7Settings.from_env()
    with db() as conn:
        message = get_message(conn, message_id)
        if not message:
            return False, "Mensaje HL7 no encontrado."
        if message["status"] != "READY":
            return False, "El mensaje debe estar READY antes de enviar."
        try:
            receipt = get_hl7_transport(settings, BASE_DIR).send(message["payload"], message["control_id"], message["message_type"])
            update_message_status(conn, message_id, "SENT")
            audit_event(conn, "hl7_sent", "hl7_message", str(message_id), user_id=user_id, encounter_id=message.get("encounter_id"), metadata={"control_id": message["control_id"], "receipt": receipt})
            if settings.mode == "mllp" and receipt:
                ack = parse_ack(receipt)
                update_message_status(conn, message_id, "ACKED" if ack.is_ack else "NACKED", ack_payload=receipt, error_message=None if ack.is_ack else ack.text)
                audit_event(conn, "hl7_ack_received" if ack.is_ack else "hl7_nack_received", "hl7_message", str(message_id), user_id=user_id, encounter_id=message.get("encounter_id"), metadata={"msa": ack.status, "control_id": ack.control_id})
            conn.commit()
            return True, receipt
        except Exception as exc:
            update_message_status(conn, message_id, "ERROR", error_message=str(exc))
            audit_event(conn, "hl7_send_failed", "hl7_message", str(message_id), user_id=user_id, encounter_id=message.get("encounter_id"), metadata={"error": str(exc)})
            conn.commit()
            return False, str(exc)


# =========================================================
# Usuarios, agenda, perfiles y correo
# =========================================================
def _is_default_admin_password(password: str) -> bool:
    return is_insecure_default_password(password)


def validate_startup_security():
    try:
        validate_admin_password_policy(APP_ENV, DEFAULT_ADMIN_PASSWORD)
    except RuntimeError:
        st.error("APP_ENV=production requiere configurar PJ_ADMIN_PASSWORD con una contraseña segura. No se permite admin1234 ni contraseña vacía.")
        st.stop()
    if SMTP_ENABLED and (not SMTP_USER or not SMTP_PASSWORD):
        st.warning("SMTP_ENABLED=true, pero SMTP_USER/SMTP_PASSWORD no están configurados. No se enviarán correos clínicos.")
    if AUTO_EMAIL_ON_NOTE:
        st.warning("AUTO_EMAIL_ON_NOTE=true. No envíes datos clínicos por correo sin autorización institucional.")


def _hash_password(password: str, salt: Optional[str] = None) -> str:
    if salt is None and PASSWORD_CONTEXT is not None:
        return PASSWORD_CONTEXT.hash(password)
    salt = salt or hashlib.sha256(os.urandom(16)).hexdigest()
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120000).hex()
    return f"{salt}${digest}"

def _check_password(password: str, stored: str) -> bool:
    if stored.startswith("$argon2") and PASSWORD_CONTEXT is not None:
        try:
            return bool(PASSWORD_CONTEXT.verify(password, stored))
        except Exception:
            return False
    try:
        salt, digest = stored.split("$", 1)
        return _hash_password(password, salt).split("$", 1)[1] == digest
    except Exception:
        return False


def _password_needs_rehash(stored: str) -> bool:
    if not stored:
        return True
    if stored.startswith("$argon2"):
        return PASSWORD_CONTEXT.needs_update(stored) if PASSWORD_CONTEXT is not None else False
    return PASSWORD_CONTEXT is not None

def ensure_default_admin_and_facility():
    ts = now_iso()
    with db() as conn:
        admin_user = conn.execute(
            "SELECT id, role, active FROM users WHERE lower(email)=lower(?)",
            (DEFAULT_ADMIN_EMAIL,),
        ).fetchone()
        if admin_user:
            if admin_user["role"] != ADMIN_ROLE or int(admin_user["active"] or 0) != 1:
                conn.execute(
                    "UPDATE users SET role=?, active=1, updated_at=? WHERE id=?",
                    (ADMIN_ROLE, ts, admin_user["id"]),
                )
        else:
            if _is_default_admin_password(DEFAULT_ADMIN_PASSWORD):
                raise RuntimeError("Configura PJ_ADMIN_PASSWORD con una contraseña inicial segura antes de crear el administrador.")
            conn.execute(
                "INSERT INTO users(email,password_hash,role,full_name,phone,specialty,preferred_note_format,first_login_completed,active,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (DEFAULT_ADMIN_EMAIL, _hash_password(DEFAULT_ADMIN_PASSWORD), ADMIN_ROLE, 'Administrador Simon Assistant', '', '', 'SOAP', 1, 1, ts, ts),
            )

        if DEFAULT_ADMIN_EMAIL != LEGACY_DEFAULT_ADMIN_EMAIL:
            legacy_admin = conn.execute(
                "SELECT id, role FROM users WHERE lower(email)=lower(?)",
                (LEGACY_DEFAULT_ADMIN_EMAIL,),
            ).fetchone()
            if legacy_admin and legacy_admin["role"] == ADMIN_ROLE:
                conn.execute(
                    "UPDATE users SET role=?, updated_at=? WHERE id=?",
                    (DOCTOR_ROLE, ts, legacy_admin["id"]),
                )
        if conn.execute("SELECT COUNT(*) AS n FROM facilities").fetchone()["n"] == 0:
            conn.execute("INSERT INTO facilities(name,code,address,phone,active,created_at,updated_at) VALUES (?,?,?,?,?,?,?)", ('Consulta ambulatoria','CA','','',1,ts,ts))
        for row in conn.execute("SELECT id,name FROM facilities WHERE code IS NULL OR trim(code)=''").fetchall():
            base_code = suggest_facility_code(row["name"])
            code = base_code
            n = 2
            while conn.execute("SELECT 1 FROM facilities WHERE code=? AND id<>?", (code, row["id"])).fetchone():
                code = f"{base_code}{n}"
                n += 1
            conn.execute("UPDATE facilities SET code=?, updated_at=? WHERE id=?", (code, ts, row["id"]))
        default_facility = conn.execute("SELECT id FROM facilities WHERE active=1 ORDER BY id LIMIT 1").fetchone()
        if default_facility:
            conn.execute(
                "UPDATE patients SET facility_id=?, updated_at=? WHERE facility_id IS NULL",
                (default_facility["id"], ts),
            )
        conn.commit()

def get_user_by_email(email: str) -> Optional[dict]:
    with db() as conn:
        row = conn.execute("SELECT * FROM users WHERE lower(email)=lower(?) AND active=1", (email.strip(),)).fetchone()
        return dict(row) if row else None

def get_user(user_id: int) -> Optional[dict]:
    with db() as conn:
        row = conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
        return dict(row) if row else None

def authenticate(email: str, password: str) -> Optional[dict]:
    user = get_user_by_email(email)
    if not user or not _check_password(password, user.get('password_hash','')):
        return None
    if _password_needs_rehash(user.get("password_hash", "")):
        with db() as conn:
            conn.execute("UPDATE users SET password_hash=?, updated_at=? WHERE id=?", (_hash_password(password), now_iso(), user["id"]))
            conn.commit()
        user = get_user(user["id"]) or user
    return user

def create_user(email: str, password: str, role: str, full_name: str, phone: str='', specialty: str=''):
    ts=now_iso()
    role = normalize_user_role(role)
    with db() as conn:
        conn.execute("INSERT INTO users(email,password_hash,role,full_name,phone,specialty,preferred_note_format,first_login_completed,active,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)", (email.strip().lower(),_hash_password(password),role,full_name,phone,specialty,'SOAP',0,1,ts,ts))
        conn.commit()

def update_current_user_profile(user_id:int, full_name:str, phone:str, specialty:str, professional_id:str, preferred_note_format:str):
    with db() as conn:
        conn.execute("UPDATE users SET full_name=?,phone=?,specialty=?,professional_id=?,preferred_note_format=?,first_login_completed=1,updated_at=? WHERE id=?", (full_name,phone,specialty,professional_id,preferred_note_format,now_iso(),user_id))
        conn.commit()

def get_users_df() -> pd.DataFrame:
    with db() as conn:
        return pd.read_sql_query("SELECT id,email,role,full_name,phone,specialty,active,updated_at,professional_id,preferred_note_format FROM users ORDER BY role, full_name", conn)

def update_user(user_id: int, email: str, full_name: str, phone: str, specialty: str, role: str, active: int, professional_id: str, preferred_note_format: str):
    email = (email or "").strip().lower()
    if not email:
        raise ValueError("El correo del usuario es obligatorio.")
    role = normalize_user_role(role)
    with db() as conn:
        conn.execute(
            "UPDATE users SET email=?, full_name=?, phone=?, specialty=?, role=?, active=?, professional_id=?, preferred_note_format=?, updated_at=? WHERE id=?",
            (email, full_name.strip(), phone.strip(), specialty.strip(), role, active, professional_id.strip(), preferred_note_format, now_iso(), user_id),
        )
        conn.commit()


def get_provider_encounters_df(provider_user_id: int) -> pd.DataFrame:
    with db() as conn:
        return pd.read_sql_query(
            """
            SELECT e.id, e.patient_id, COALESCE(p.mrn,'') AS mrn,
                   TRIM(COALESCE(p.first_name,'') || ' ' || COALESCE(p.last_name,'')) AS patient_name,
                   e.provider_name, e.facility_name, e.visit_type, e.chief_complaint,
                   e.status, e.created_at, e.updated_at
            FROM encounters e
            JOIN appointments a ON a.encounter_id = e.id
            LEFT JOIN patients p ON p.id = e.patient_id
            WHERE a.provider_user_id = ?
            ORDER BY e.created_at DESC
            """,
            conn,
            params=(provider_user_id,),
        )

def get_facilities_df(active_only: bool=True) -> pd.DataFrame:
    with db() as conn:
        q="SELECT id,name,code,address,phone,active FROM facilities" + (" WHERE active=1" if active_only else "") + " ORDER BY name"
        return pd.read_sql_query(q, conn)

def create_facility(name:str, code:str, address:str='', phone:str=''):
    code = normalize_facility_code(code)
    ts=now_iso()
    with db() as conn:
        if conn.execute("SELECT 1 FROM facilities WHERE code=?", (code,)).fetchone():
            raise ValueError("Ya existe un centro con esa sigla.")
        conn.execute("INSERT INTO facilities(name,code,address,phone,active,created_at,updated_at) VALUES (?,?,?,?,?,?,?)", (name.strip(),code,address.strip(),phone.strip(),1,ts,ts))
        conn.commit()

def update_facility(facility_id:int, name:str, code:str, address:str='', phone:str='', active:int=1):
    code = normalize_facility_code(code)
    if not name.strip():
        raise ValueError("El nombre del centro es obligatorio.")
    with db() as conn:
        if conn.execute("SELECT 1 FROM facilities WHERE code=? AND id<>?", (code, facility_id)).fetchone():
            raise ValueError("Ya existe otro centro con esa sigla.")
        conn.execute(
            "UPDATE facilities SET name=?, code=?, address=?, phone=?, active=?, updated_at=? WHERE id=?",
            (name.strip(), code, address.strip(), phone.strip(), active, now_iso(), facility_id),
        )
        conn.commit()

def create_appointment(patient_id, facility_id:int, provider_user_id:int, appointment_date:str, appointment_time:str, duration:int, reason:str):
    ts=now_iso()
    with db() as conn:
        patient = conn.execute("SELECT facility_id FROM patients WHERE id=?", (patient_id,)).fetchone()
        if not patient:
            raise ValueError("Debes seleccionar un paciente válido.")
        if int(patient["facility_id"] or 0) != int(facility_id or 0):
            raise ValueError("El paciente seleccionado no pertenece al centro de la cita.")
        cur = conn.execute("INSERT INTO appointments(patient_id,facility_id,provider_user_id,appointment_date,appointment_time,duration_minutes,status,reason,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?)", (patient_id,facility_id,provider_user_id,appointment_date,appointment_time,duration,'agendada',reason,ts,ts))
        conn.commit()
        return cur.lastrowid

def get_appointments_df(day:str='', facility_id=None, provider_user_id=None) -> pd.DataFrame:
    clauses=[]; params=[]
    if day: clauses.append('a.appointment_date=?'); params.append(day)
    if facility_id: clauses.append('a.facility_id=?'); params.append(facility_id)
    if provider_user_id: clauses.append('a.provider_user_id=?'); params.append(provider_user_id)
    where=(' WHERE '+ ' AND '.join(clauses)) if clauses else ''
    with db() as conn:
        return pd.read_sql_query(f"""SELECT a.id,a.appointment_date AS fecha,a.appointment_time AS hora,a.duration_minutes AS min,a.status AS estado,f.name AS centro,COALESCE(u.full_name,u.email,'') AS profesional,COALESCE(p.first_name||' '||p.last_name,'Paciente no asignado') AS paciente,p.mrn,a.reason AS motivo,a.patient_id,a.encounter_id FROM appointments a LEFT JOIN patients p ON p.id=a.patient_id LEFT JOIN facilities f ON f.id=a.facility_id LEFT JOIN users u ON u.id=a.provider_user_id {where} ORDER BY f.name,a.appointment_date,a.appointment_time""", conn, params=params)

def build_daily_slots_df(day_value: date, facilities_df: pd.DataFrame, providers_df: pd.DataFrame, appts_df: pd.DataFrame, start_time: str, end_time: str, slot_minutes: int) -> pd.DataFrame:
    if facilities_df.empty or providers_df.empty:
        return pd.DataFrame()
    day_text = day_value.isoformat()
    start_dt = datetime.strptime(f"{day_text} {start_time}", "%Y-%m-%d %H:%M")
    end_dt = datetime.strptime(f"{day_text} {end_time}", "%Y-%m-%d %H:%M")
    if end_dt <= start_dt:
        return pd.DataFrame()

    rows = []
    cursor = start_dt
    slots = []
    while cursor < end_dt:
        slot_end = cursor + pd.Timedelta(minutes=slot_minutes)
        slots.append((cursor, min(slot_end, end_dt)))
        cursor = slot_end

    for _, facility in facilities_df.iterrows():
        for _, provider in providers_df.iterrows():
            facility_name = facility["name"]
            provider_name = provider.get("full_name")
            if pd.isna(provider_name) or not str(provider_name).strip():
                provider_name = provider.get("email") or ""
            provider_name = str(provider_name)
            provider_appts = appts_df[(appts_df["centro"] == facility_name) & (appts_df["profesional"] == provider_name)] if not appts_df.empty else pd.DataFrame()
            for slot_start, slot_end in slots:
                match = None
                if not provider_appts.empty:
                    for _, appt in provider_appts.iterrows():
                        appt_start = datetime.strptime(f"{appt['fecha']} {appt['hora']}", "%Y-%m-%d %H:%M")
                        appt_end = appt_start + pd.Timedelta(minutes=int(appt.get("min") or slot_minutes))
                        if slot_start < appt_end and slot_end > appt_start and appt.get("estado") != "cancelada":
                            match = appt
                            break
                rows.append({
                    "Centro": facility_name,
                    "Profesional": provider_name,
                    "Hora": slot_start.strftime("%H:%M"),
                    "Cupo": "Tomado" if match is not None else "Disponible",
                    "Paciente": "" if match is None else match.get("paciente", ""),
                    "Estado cita": "" if match is None else match.get("estado", ""),
                    "ID cita": "" if match is None else int(match.get("id")),
                    "_fecha": day_text,
                    "_facility_id": int(facility["id"]),
                    "_provider_user_id": int(provider["id"]),
                    "_duration_minutes": slot_minutes,
                    "_patient_id": "" if match is None else int(match.get("patient_id") or 0),
                    "_encounter_id": "" if match is None or pd.isna(match.get("encounter_id")) else int(match.get("encounter_id") or 0),
                })
    return pd.DataFrame(rows)

def set_appointment_status(appt_id:int, status:str):
    with db() as conn:
        conn.execute("UPDATE appointments SET status=?,updated_at=? WHERE id=?", (status,now_iso(),appt_id)); conn.commit()

def attach_encounter_to_appointment(appt_id:int, encounter_id:int):
    with db() as conn:
        conn.execute("UPDATE appointments SET encounter_id=?,updated_at=? WHERE id=?", (encounter_id,now_iso(),appt_id)); conn.commit()

def get_patient_appointments_on_day(patient_id:int, day_text:str) -> pd.DataFrame:
    with db() as conn:
        return pd.read_sql_query(
            """
            SELECT a.id, a.appointment_time AS hora, a.duration_minutes AS min, a.status AS estado,
                   f.name AS centro, COALESCE(u.full_name,u.email,'') AS profesional, a.reason AS motivo
            FROM appointments a
            LEFT JOIN facilities f ON f.id=a.facility_id
            LEFT JOIN users u ON u.id=a.provider_user_id
            WHERE a.patient_id=? AND a.appointment_date=? AND a.status <> 'cancelada'
            ORDER BY a.appointment_time
            """,
            conn,
            params=(patient_id, day_text),
        )

def schedule_or_confirm_same_day(patient_id:int, facility_id:int, provider_user_id:int, appointment_date:str, appointment_time:str, duration:int, reason:str) -> bool:
    same_day = get_patient_appointments_on_day(patient_id, appointment_date)
    payload = {
        'patient_id': patient_id,
        'facility_id': facility_id,
        'provider_user_id': provider_user_id,
        'appointment_date': appointment_date,
        'appointment_time': appointment_time,
        'duration_minutes': int(duration),
        'reason': reason,
    }
    if same_day.empty:
        create_appointment(patient_id, facility_id, provider_user_id, appointment_date, appointment_time, int(duration), reason)
        return True
    st.session_state['pending_same_day_schedule'] = payload
    return False

@st.dialog("Paciente con cita el mismo día")
def confirm_same_day_schedule_dialog():
    pending = st.session_state.get('pending_same_day_schedule')
    if not pending:
        return
    patient = get_patient(int(pending['patient_id']))
    same_day = get_patient_appointments_on_day(int(pending['patient_id']), pending['appointment_date'])
    st.warning("Este paciente ya tiene una cita el mismo día.")
    if patient:
        st.write(f"**Paciente:** {patient.get('mrn')} - {patient.get('first_name','')} {patient.get('last_name','')}")
    st.dataframe(same_day.rename(columns={'hora':'Hora','min':'Min','estado':'Estado','centro':'Centro','profesional':'Profesional','motivo':'Motivo'}), use_container_width=True, hide_index=True)
    st.write(f"**Nuevo cupo:** {pending['appointment_date']} {pending['appointment_time']}")
    cancel_col, confirm_col = st.columns(2)
    with cancel_col:
        if st.button("Cancelar agendamiento"):
            st.session_state['pending_same_day_schedule'] = None
            st.rerun()
    with confirm_col:
        if st.button("Agendar de todas formas"):
            try:
                create_appointment(
                    int(pending['patient_id']),
                    int(pending['facility_id']),
                    int(pending['provider_user_id']),
                    pending['appointment_date'],
                    pending['appointment_time'],
                    int(pending['duration_minutes']),
                    pending.get('reason', ''),
                )
                st.session_state['pending_same_day_schedule'] = None
                st.session_state['agenda_message'] = 'Cita guardada con advertencia de cita previa el mismo día.'
                st.rerun()
            except Exception as e:
                st.error(f'No se pudo guardar la cita: {e}')

def start_clinical_from_appointment(appt_row, user:dict) -> Optional[int]:
    patient_id_raw = appt_row.get('patient_id')
    if pd.isna(patient_id_raw) or not int(patient_id_raw or 0):
        return None

    patient_id = int(patient_id_raw)
    encounter_raw = appt_row.get('encounter_id')
    encounter_id = None if pd.isna(encounter_raw) or not int(encounter_raw or 0) else int(encounter_raw)
    if not encounter_id:
        encounter_id = create_encounter(patient_id, {
            "visit_type": "INGRESO",
            "brought_by": "",
            "informant": "",
            "provider_name": appt_row.get('profesional') or user.get('full_name') or user.get('email') or "",
            "facility_name": appt_row.get('centro') or "",
            "chief_complaint": appt_row.get('motivo') or "",
        })
        attach_encounter_to_appointment(int(appt_row.get('id')), encounter_id)

    st.session_state['patient_id'] = patient_id
    patient = get_patient(patient_id)
    st.session_state['patient_facility_id'] = patient.get('facility_id') if patient else None
    st.session_state['encounter_id'] = encounter_id
    st.session_state['appointment_id'] = int(appt_row.get('id'))
    st.session_state['active_appointment_id'] = int(appt_row.get('id'))
    set_appointment_status(int(appt_row.get('id')), 'en_atencion')
    st.session_state['active_page'] = 'Atención clínica'
    return encounter_id

def close_clinical_attention(encounter_id:int, message:str='Atención cerrada.'):
    update_encounter_content(encounter_id, status='completed')
    appt_id = st.session_state.get('active_appointment_id')
    if not appt_id:
        with db() as conn:
            row = conn.execute("SELECT id FROM appointments WHERE encounter_id=? ORDER BY id DESC LIMIT 1", (encounter_id,)).fetchone()
            appt_id = row["id"] if row else None
    if appt_id:
        set_appointment_status(int(appt_id), 'atendida')
    for key in ['patient_id', 'encounter_id', 'appointment_id', 'active_appointment_id']:
        st.session_state[key] = None
    st.session_state['agenda_message'] = message
    st.session_state['active_page'] = 'Agenda'
    st.session_state.pop("reviewed_note_text", None)
    st.session_state.pop("reviewed_note_encounter_id", None)

@st.dialog("Confirmar envío de nota clínica")
def confirm_reviewed_note_send_dialog():
    pending = st.session_state.get("pending_note_send")
    if not pending:
        return
    to_email = (pending.get("to_email") or "").strip()
    encounter_id = int(pending["encounter_id"])
    reviewed_note = pending.get("note_text") or ""
    st.warning("Se enviará la nota clínica revisada por correo y se cerrará la atención.")
    st.write(f"**Destino:** {to_email or 'sin correo configurado'}")
    st.caption("Si cancelas, no se enviará nada y podrás seguir editando la nota.")
    cancel_col, send_col = st.columns(2)
    with cancel_col:
        if st.button("Cancelar y seguir editando"):
            st.session_state["pending_note_send"] = None
            st.rerun()
    with send_col:
        if st.button("Enviar y cerrar atención"):
            if not to_email:
                st.error("El profesional conectado no tiene correo configurado.")
                return
            update_encounter_content(encounter_id, note_text=reviewed_note)
            out = export_docx("Ficha clínica ambulatoria", reviewed_note, f"nota_revisada_encuentro_{encounter_id}_{int(time.time())}.docx")
            ok, msg = send_email_with_attachment(
                to_email,
                f"Simon Assistant - Nota clínica revisada encuentro {encounter_id}",
                "Se adjunta la nota clínica revisada desde Simon Assistant.",
                out,
            )
            if ok:
                with db() as conn:
                    audit_event(conn, "email_sent", "encounter", str(encounter_id), user_id=(st.session_state.get("auth_user") or {}).get("id"), encounter_id=encounter_id, metadata={"to": to_email, "attachment": out})
                    conn.commit()
                st.session_state["pending_note_send"] = None
                close_clinical_attention(encounter_id, f"Nota revisada enviada a {to_email}. Atención marcada como atendida y cerrada.")
                st.rerun()
            else:
                st.error(msg)
                st.info(f"Word generado localmente: {out}")

def render_soap_from_note(note_text:str, note_json:Optional[dict], meta:dict) -> str:
    if not note_json:
        return "REGISTRO CLINICO FORMATO SOAP\n\n" + (note_text or "")
    def line(label,key):
        v=fv(note_json,key)
        return f"- {label}: {v}" if v else ""
    subj=[line("Motivo","chief_complaint"),line("Inicio","hpi_onset"),line("Curso","hpi_course"),line("Frecuencia","hpi_frequency"),line("Antecedentes","past_medical_history")]
    obj=[line("Examen neurológico","exam_neuro"),line("Exámenes revisados","exams_reviewed")]
    ana=[line("Hipótesis diagnóstica","diagnostic_hypothesis")]
    plan=[line("Generales","indications.general"),line("Colegio","indications.school"),line("Medicamentos","indications.meds"),line("Control","indications.followup"),line("Derivaciones","indications.referrals")]
    clean=lambda xs: "\n".join([x for x in xs if x]) or "No mencionado"
    header = f"REGISTRO CLINICO FORMATO SOAP\n\nPaciente: {meta.get('patient_name','')}\nProfesional: {meta.get('provider_name','')}\nCentro: {meta.get('facility_name','')}\n"
    return header + f"\nS - Subjetivo:\n{clean(subj)}\n\nO - Objetivo:\n{clean(obj)}\n\nA - Evaluación / Análisis:\n{clean(ana)}\n\nP - Plan:\n{clean(plan)}\n"

def get_smtp_config() -> dict:
    smtp_user = (get_setting("SMTP_USER", "") or "").strip()
    return {
        "enabled": get_bool_setting("SMTP_ENABLED", False),
        "host": get_setting("SMTP_HOST", "smtp.gmail.com").strip(),
        "port": get_int_setting("SMTP_PORT", 587),
        "user": smtp_user,
        "password": get_setting("SMTP_PASSWORD", "").replace(" ", "").strip(),
        "from": get_setting("SMTP_FROM", smtp_user).strip(),
    }

def mask_secret(value:str) -> str:
    clean = (value or "").replace(" ", "").strip()
    if not clean:
        return "(no configurado)"
    return f"{clean[:4]}{'*' * max(len(clean) - 8, 4)}{clean[-4:]}"

def send_email_with_attachment(to_email:str, subject:str, body:str, attachment_path:str) -> Tuple[bool,str]:
    cfg = get_smtp_config()
    to_email = (to_email or "").strip()
    if not to_email:
        return False, 'El profesional autenticado no tiene correo configurado.'
    if not cfg["enabled"]:
        return False, 'SMTP deshabilitado. Configura SMTP_ENABLED=true solo con autorización institucional para enviar datos clínicos.'
    if not cfg["user"] or not cfg["password"]:
        return False, 'Correo no configurado. Completa SMTP_USER y SMTP_PASSWORD como variables de entorno o en st.secrets. Para Gmail usa App Password.'
    msg=EmailMessage(); msg['From']=cfg["from"]; msg['To']=to_email; msg['Subject']=subject; msg.set_content(body)
    with open(attachment_path,'rb') as f: data=f.read()
    msg.add_attachment(data, maintype='application', subtype='vnd.openxmlformats-officedocument.wordprocessingml.document', filename=os.path.basename(attachment_path))
    try:
        with smtplib.SMTP(cfg["host"], cfg["port"], timeout=30) as server:
            server.starttls(context=ssl.create_default_context()); server.login(cfg["user"], cfg["password"]); server.send_message(msg)
        return True, 'Correo enviado correctamente.'
    except Exception as e:
        return False, f'No se pudo enviar el correo: {e}'

def env_flag(name: str, default: bool = False) -> bool:
    return get_bool_setting(name, default)


def valid_env_secret(value: Optional[str]) -> bool:
    value = (value or "").strip()
    if not value:
        return False
    lowered = value.lower()
    return not (
        lowered.startswith("tu_") or
        lowered in {"changeme", "change_me", "pendiente", "xxx", "..."}
    )


def google_login_enabled() -> bool:
    return env_flag("GOOGLE_LOGIN_ENABLED", False)


def google_oauth_configured() -> bool:
    return google_login_enabled() and all([
        valid_env_secret(get_setting('GOOGLE_CLIENT_ID')),
        valid_env_secret(get_setting('GOOGLE_CLIENT_SECRET')),
        valid_env_secret(get_setting('GOOGLE_REDIRECT_URI')),
    ])


def google_oauth_authorization_url() -> Optional[str]:
    if not google_oauth_configured():
        return None
    params = {
        'client_id': get_setting('GOOGLE_CLIENT_ID'),
        'redirect_uri': get_setting('GOOGLE_REDIRECT_URI'),
        'response_type': 'code',
        'scope': 'openid email profile',
        'access_type': 'offline',
        'prompt': 'select_account',
        'state': str(uuid.uuid4()),
    }
    return 'https://accounts.google.com/o/oauth2/v2/auth?' + urllib.parse.urlencode(params)


def exchange_google_code(code: str) -> Optional[dict]:
    token_url = 'https://oauth2.googleapis.com/token'
    data = {
        'code': code,
        'client_id': get_setting('GOOGLE_CLIENT_ID'),
        'client_secret': get_setting('GOOGLE_CLIENT_SECRET'),
        'redirect_uri': get_setting('GOOGLE_REDIRECT_URI'),
        'grant_type': 'authorization_code',
    }
    try:
        resp = requests.post(token_url, data=data, timeout=15)
        resp.raise_for_status()
        return resp.json()
    except Exception:
        return None


def fetch_google_userinfo(access_token: str) -> Optional[dict]:
    try:
        resp = requests.get(
            'https://openidconnect.googleapis.com/v1/userinfo',
            headers={'Authorization': f'Bearer {access_token}'},
            timeout=15,
        )
        resp.raise_for_status()
        return resp.json()
    except Exception:
        return None


def handle_google_callback() -> Optional[dict]:
    params = st.query_params
    if 'code' not in params:
        return None
    code = params.get('code', [''])[0]
    token_data = exchange_google_code(code)
    if not token_data or not token_data.get('access_token'):
        st.error('No se pudo completar el inicio de sesión con Google. Verifica las credenciales de OAuth.')
        return None
    userinfo = fetch_google_userinfo(token_data['access_token'])
    if not userinfo or not userinfo.get('email'):
        st.error('No se pudo obtener información de la cuenta Google.')
        return None
    email = userinfo['email'].strip().lower()
    user = get_user_by_email(email)
    if not user:
        st.error('Esta cuenta Google no está autorizada. Solicita al administrador que habilite tu correo.')
        return None
    try:
        st.query_params.clear()
    except Exception:
        pass
    return user


def require_login():
    try:
        ensure_default_admin_and_facility()
    except RuntimeError as exc:
        st.error(str(exc))
        st.stop()
    st.session_state.setdefault('auth_user', None)
    if st.session_state.get('auth_user'):
        cached_user = st.session_state['auth_user']
        refreshed_user = get_user(cached_user.get('id')) if cached_user.get('id') else get_user_by_email(cached_user.get('email', ''))
        if refreshed_user:
            st.session_state['auth_user'] = refreshed_user
            return refreshed_user
        st.session_state['auth_user'] = None
    google_user = handle_google_callback()
    if google_user:
        st.session_state['auth_user'] = google_user
        st.rerun()
    left, center, right = st.columns([1, 1.15, 1])
    with center:
        render_app_brand()
        render_page_header(APP_NAME, 'Utilice las credenciales otorgadas por el administrador.')
        with st.container(border=True):
            with st.form('login_form'):
                email = st.text_input('Correo', value=DEFAULT_ADMIN_EMAIL)
                password = st.text_input('Contraseña', type='password')
                submitted = st.form_submit_button('Ingresar')
            if submitted:
                user = authenticate(email, password)
                if user:
                    st.session_state['auth_user'] = user
                    st.rerun()
                else:
                    st.error('Correo o contraseña incorrectos, o usuario inactivo.')
            google_url = google_oauth_authorization_url()
            if google_url:
                st.divider()
                st.markdown(
                    f'''<a href="{google_url}" style="display:inline-flex; align-items:center; justify-content:center; gap:0.75rem; width:100%; padding:0.85rem 1rem; border-radius:999px; border:1px solid #dadce0; background:#fff; color:#202124; text-decoration:none; font-weight:600; box-shadow:0 1px 2px rgba(0,0,0,0.05);">
                            <span style="display:inline-flex; align-items:center; justify-content:center; width:18px; height:18px; color:#4285f4; font-weight:800; font-family:Arial, sans-serif;">G</span>
                            Continuar con Google
                        </a>''',
                    unsafe_allow_html=True,
                )
                st.caption('Inicia sesión con tu cuenta de Google si ya la tienes configurada en la consola de Google Cloud.')
            elif google_login_enabled():
                st.divider()
                st.warning('Para habilitar ingreso con Google, configura GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET y GOOGLE_REDIRECT_URI como variables de entorno o en st.secrets.')
    st.stop()

def require_profile_completion(user:dict):
    if user.get('first_login_completed'): return user
    render_page_header('Completar perfil profesional', 'Primer ingreso: completa tus datos. Este correo recibirá automáticamente una copia Word del registro clínico.')
    with st.form('profile_form'):
        full_name=st.text_input('Nombre completo', value=user.get('full_name') or '')
        phone=st.text_input('Número de contacto', value=user.get('phone') or '')
        specialty=st.text_input('Especialidad / rol', value=user.get('specialty') or '')
        professional_id=st.text_input('Registro profesional / identificador interno', value=user.get('professional_id') or '')
        preferred=st.selectbox('Formato preferido de envío automático', ['SOAP','Ficha clínica completa'])
        ok=st.form_submit_button('Guardar perfil')
    if ok:
        if not full_name.strip() or not phone.strip(): st.error('Nombre completo y teléfono son obligatorios.')
        else:
            update_current_user_profile(user['id'],full_name,phone,specialty,professional_id,preferred); st.session_state['auth_user']=get_user(user['id']); st.rerun()
    st.stop()


def render_main_menu(user: dict):
    render_page_header(APP_NAME, f"Bienvenido, {user.get('full_name') or user.get('email')}")
    st.write('Selecciona el módulo que quieres usar:')
    cols = st.columns(4)
    with cols[0]:
        render_nav_card('Agenda', 'Ver y gestionar citas médicas rápidas.', 'Agenda', 'Agenda', 'menu_agenda')
    with cols[1]:
        render_nav_card('Atención clínica', 'Buscar paciente, abrir ficha y continuar la consulta.', 'Atención clínica', 'Atención clínica', 'menu_atencion_clinica')
    with cols[2]:
        render_nav_card('Historial atenciones de pacientes', 'Buscar pacientes y revisar su ficha e historial de atenciones.', 'Historial atenciones de pacientes', 'Pacientes', 'menu_pacientes')
    if is_admin_user(user):
        with cols[3]:
            render_nav_card('Administrador', 'Gestionar profesionales, usuarios y configuración.', 'Administrador', 'Administrador', 'menu_administrador')
    st.markdown('---')
    st.write('Usa el botón de cerrar sesión en la barra lateral para salir.')
    st.stop()


def render_agenda_page(user:dict):
    render_page_header('Agenda ambulatoria', 'Gestiona citas, pacientes, centros y el inicio de consultas desde una vista simple.')
    if st.session_state.get('agenda_message'):
        st.success(st.session_state.pop('agenda_message'))
    if st.session_state.get('pending_same_day_schedule'):
        confirm_same_day_schedule_dialog()
    facilities=get_facilities_df(); users=get_users_df(); patients=get_patients_df(); doctors=users[users['role'].isin([DOCTOR_ROLE, ADMIN_ROLE])] if not users.empty else users
    with st.container(border=True):
        st.markdown('**Filtros de agenda**')
        c1,c2,c3=st.columns(3)
        with c1: day=st.date_input('Día', value=date.today(), min_value=date(1970,1,1), format='DD/MM/YYYY')
        with c2: facility_label=st.selectbox('Centro', ['Todos']+facilities['name'].tolist() if not facilities.empty else ['Todos'])
        with c3:
            if user.get('role')==DOCTOR_ROLE: provider_id=user['id']; st.text_input('Profesional', value=user.get('full_name') or user.get('email'), disabled=True)
            else:
                opts=['Todos']+doctors.apply(lambda r:f"{r['id']} - {r['full_name'] or r['email']}",axis=1).tolist() if not doctors.empty else ['Todos']
                provider_label=st.selectbox('Profesional',opts); provider_id=None if provider_label=='Todos' else int(provider_label.split(' - ')[0])
    facility_id=None if facility_label=='Todos' or facilities.empty else int(facilities.loc[facilities['name']==facility_label,'id'].iloc[0])
    appts=get_appointments_df(day.isoformat(), facility_id, provider_id)
    selected_appt_row = None
    selected_slot_row = None

    with st.expander('Agenda completa del día', expanded=True):
        s1,s2,s3=st.columns(3)
        with s1:
            agenda_start=st.time_input('Inicio jornada', value=dt_time(8,0), key='agenda_full_start')
        with s2:
            agenda_end=st.time_input('Fin jornada', value=dt_time(18,0), key='agenda_full_end')
        with s3:
            slot_minutes=st.selectbox('Duración cupo', [15,20,30,45,60], index=2, key='agenda_full_slot')
        slot_facilities = facilities[facilities['id'] == facility_id] if facility_id and not facilities.empty else facilities
        slot_providers = doctors[doctors['id'] == provider_id] if provider_id and not doctors.empty else doctors
        slots_df = build_daily_slots_df(
            day,
            slot_facilities,
            slot_providers,
            appts,
            agenda_start.strftime('%H:%M'),
            agenda_end.strftime('%H:%M'),
            int(slot_minutes),
        )
        if slots_df.empty:
            st.info('No hay centros o profesionales para construir la agenda diaria.')
        else:
            taken = int((slots_df['Cupo'] == 'Tomado').sum())
            available = int((slots_df['Cupo'] == 'Disponible').sum())
            m1,m2=st.columns(2)
            with m1:
                render_metric_card('Cupos disponibles', available, 'Cupos libres según filtros actuales')
            with m2:
                render_metric_card('Cupos tomados', taken, 'Citas agendadas no canceladas')
            hidden_cols = ['_fecha','_facility_id','_provider_user_id','_duration_minutes','_patient_id','_encounter_id']
            visible_slots = slots_df.drop(columns=[c for c in hidden_cols if c in slots_df.columns])
            slot_selection = st.dataframe(
                visible_slots,
                use_container_width=True,
                hide_index=True,
                on_select='rerun',
                selection_mode='single-row',
                key='agenda_full_action_table',
            )
            selected_rows = slot_selection.selection.rows if slot_selection and slot_selection.selection else []
            if selected_rows:
                selected_slot_row = slots_df.iloc[int(selected_rows[0])]
            if selected_slot_row is not None:
                if selected_slot_row.get('Cupo') == 'Tomado' and selected_slot_row.get('ID cita'):
                    row = appts[appts['id'] == int(selected_slot_row['ID cita'])]
                    if not row.empty:
                        selected_appt_row = row.iloc[0]

            st.markdown('**Acciones del cupo seleccionado**')
            if selected_slot_row is None:
                st.caption('Selecciona un cupo disponible para agendar o un cupo tomado para acciones de cita/paciente.')
            elif selected_appt_row is not None:
                st.caption(f"Cita #{int(selected_appt_row['id'])} - {selected_appt_row.get('paciente','Sin paciente')} - {selected_appt_row.get('hora','')}")
                x,y,z,w=st.columns(4)
                with x:
                    if st.button('Marcar llegada'):
                        set_appointment_status(int(selected_appt_row['id']),'llegó'); st.success('Cita actualizada.'); st.rerun()
                with y:
                    if st.button('Cancelar cita'):
                        set_appointment_status(int(selected_appt_row['id']),'cancelada'); st.success('Cita cancelada.'); st.rerun()
                with z:
                    if st.button('Atender y grabar con bot'):
                        encounter_id = start_clinical_from_appointment(selected_appt_row, user)
                        if encounter_id:
                            st.rerun()
                        else:
                            st.error('La cita no tiene paciente asociado.')
                with w:
                    patient_id_from_appt = int(selected_appt_row.get('patient_id') or 0)
                    if st.button('Abrir ficha paciente', disabled=not patient_id_from_appt):
                        patient = get_patient(patient_id_from_appt)
                        st.session_state['patient_id'] = patient_id_from_appt
                        st.session_state['patient_facility_id'] = patient.get('facility_id') if patient else None
                        st.session_state['active_page'] = 'Pacientes'
                        st.rerun()
            else:
                st.caption(f"Cupo disponible: {selected_slot_row['Centro']} - {selected_slot_row['Profesional']} - {selected_slot_row['Hora']}")
                selected_facility_id = int(selected_slot_row['_facility_id'])
                selected_provider_id = int(selected_slot_row['_provider_user_id'])
                facility_patients = patients[patients['facility_id'] == selected_facility_id].copy() if not patients.empty else pd.DataFrame()
                tab_existing, tab_new = st.tabs(['Paciente existente', 'Paciente nuevo'])
                with tab_existing:
                    patient_search = st.text_input('Buscar paciente del centro', value='', key='slot_patient_search')
                    if patient_search.strip() and not facility_patients.empty:
                        q = patient_search.strip().lower()
                        facility_patients = facility_patients[
                            facility_patients.apply(lambda r: q in " ".join([
                                str(r.get('mrn') or ''),
                                str(r.get('national_id') or ''),
                                str(r.get('first_name') or ''),
                                str(r.get('last_name') or ''),
                            ]).lower(), axis=1)
                        ]
                    if facility_patients.empty:
                        st.info('No hay pacientes existentes para este centro.')
                    else:
                        quick_patients = facility_patients.assign(paciente=facility_patients['first_name'].fillna('') + ' ' + facility_patients['last_name'].fillna('')).reset_index(drop=True)
                        page_key = f"slot_patient_page_{selected_facility_id}"
                        st.session_state.setdefault(page_key, 0)
                        total_patients = len(quick_patients)
                        total_pages = max(1, (total_patients + 9) // 10)
                        if st.session_state[page_key] >= total_pages:
                            st.session_state[page_key] = 0
                        start_idx = st.session_state[page_key] * 10
                        end_idx = min(start_idx + 10, total_patients)
                        page_patients = quick_patients.iloc[start_idx:end_idx].reset_index(drop=True)
                        st.caption(f"Mostrando {start_idx + 1}-{end_idx} de {total_patients}")
                        patient_selection = st.dataframe(
                            page_patients[['mrn','national_id','paciente','birth_date']].rename(columns={'mrn':'ID','national_id':'RUN','paciente':'Paciente','birth_date':'Nacimiento'}),
                            use_container_width=True,
                            hide_index=True,
                            on_select='rerun',
                            selection_mode='single-row',
                            key='slot_existing_patient_table',
                        )
                        selected_patient_rows = patient_selection.selection.rows if patient_selection and patient_selection.selection else []
                        selected_patient_id = int(page_patients.iloc[int(selected_patient_rows[0])]['id']) if selected_patient_rows else None
                        nav_prev, nav_next = st.columns([1,1])
                        with nav_prev:
                            if st.button('Anterior', disabled=st.session_state[page_key] <= 0, key='slot_patient_prev'):
                                st.session_state[page_key] -= 1
                                st.rerun()
                        with nav_next:
                            if st.button('Next', disabled=st.session_state[page_key] >= total_pages - 1, key='slot_patient_next'):
                                st.session_state[page_key] += 1
                                st.rerun()
                        slot_reason = st.text_input('Motivo / observación', value='', key='slot_existing_reason')
                        if st.button('Agendar paciente seleccionado en este cupo', disabled=selected_patient_id is None):
                            try:
                                scheduled = schedule_or_confirm_same_day(selected_patient_id, selected_facility_id, selected_provider_id, selected_slot_row['_fecha'], selected_slot_row['Hora'], int(selected_slot_row['_duration_minutes']), slot_reason)
                                if scheduled:
                                    st.success('Cita guardada en el cupo seleccionado.')
                                    st.rerun()
                                else:
                                    st.rerun()
                            except Exception as e:
                                st.error(f'No se pudo guardar la cita: {e}')
                with tab_new:
                    next_mrn = preview_next_patient_mrn(selected_facility_id)
                    st.text_input('Identificador automático / MRN', value=next_mrn, disabled=True, key='slot_new_mrn_preview')
                    n1,n2,n3 = st.columns(3)
                    with n1:
                        new_first_name = st.text_input('Nombre', key='slot_new_first_name')
                        new_last_name = st.text_input('Apellidos', key='slot_new_last_name')
                    with n2:
                        new_national_id = st.text_input('RUN / documento', key='slot_new_national_id')
                        new_birth_date = st.date_input('Fecha de nacimiento', value=None, min_value=date(1900,1,1), max_value=date.today(), format='DD/MM/YYYY', key='slot_new_birth_date')
                        new_sex = st.selectbox('Sexo', ["", "M", "F", "O", "U"], key='slot_new_sex')
                    with n3:
                        new_phone = st.text_input('Teléfono', key='slot_new_phone')
                        new_email = st.text_input('Email', key='slot_new_email')
                        new_guardian = st.text_input('Adulto responsable / tutor', key='slot_new_guardian')
                    new_address = st.text_input('Dirección', key='slot_new_address')
                    new_reason = st.text_input('Motivo / observación', value='', key='slot_new_reason')
                    if st.button('Crear paciente y agendar en este cupo'):
                        if not new_first_name.strip() or not new_last_name.strip():
                            st.error('Nombre y apellidos son obligatorios.')
                        else:
                            try:
                                new_patient_id = create_or_update_patient({
                                    "mrn": None,
                                    "first_name": new_first_name,
                                    "last_name": new_last_name,
                                    "national_id": new_national_id,
                                    "birth_date": new_birth_date.isoformat() if new_birth_date else "",
                                    "sex": new_sex,
                                    "phone": new_phone,
                                    "email": new_email,
                                    "address": new_address,
                                    "guardian_name": new_guardian,
                                    "facility_id": selected_facility_id,
                                })
                                create_appointment(new_patient_id, selected_facility_id, selected_provider_id, selected_slot_row['_fecha'], selected_slot_row['Hora'], int(selected_slot_row['_duration_minutes']), new_reason)
                                saved_patient = get_patient(new_patient_id)
                                st.success(f"Paciente {saved_patient.get('mrn') if saved_patient else new_patient_id} creado y agendado.")
                                st.rerun()
                            except Exception as e:
                                st.error(f'No se pudo crear y agendar el paciente: {e}')
def render_patient_history_detail(patient_id: int):
    patient = get_patient(patient_id)
    if not patient:
        st.warning('Paciente no encontrado.')
        return

    st.markdown('### Datos del paciente')
    d1,d2,d3=st.columns(3)
    with d1:
        st.write(f"**ID/MRN:** {patient.get('mrn') or ''}")
        st.write(f"**Nombre:** {patient.get('first_name','')} {patient.get('last_name','')}")
        st.write(f"**Centro:** {patient.get('facility_name') or 'Sin centro'}")
    with d2:
        st.write(f"**RUN/documento:** {patient.get('national_id') or ''}")
        st.write(f"**Nacimiento:** {patient.get('birth_date') or ''}")
        st.write(f"**Sexo:** {patient.get('sex') or ''}")
    with d3:
        st.write(f"**Teléfono:** {patient.get('phone') or ''}")
        st.write(f"**Email:** {patient.get('email') or ''}")
        st.write(f"**Tutor/responsable:** {patient.get('guardian_name') or ''}")
    st.write(f"**Dirección:** {patient.get('address') or ''}")

    notes_df = get_patient_notes_df(patient_id)
    st.markdown('### Historial de atenciones')
    if notes_df.empty:
        st.info('No hay resúmenes ni notas creadas para este paciente.')
        return

    summary_df = notes_df[['id','created_at','visit_type','chief_complaint','status','updated_at']].rename(columns={
        'id':'ID','created_at':'Fecha','visit_type':'Tipo','chief_complaint':'Motivo','status':'Estado','updated_at':'Actualizado'
    })
    st.dataframe(summary_df, use_container_width=True, hide_index=True)

    note_options = notes_df.apply(lambda r: f"{r['id']} | {r['created_at']} | {r.get('chief_complaint') or 'Sin motivo'}", axis=1).tolist()
    selected_note = st.selectbox('Seleccionar nota para visualizar', [''] + note_options, key=f'patient_note_select_{patient_id}')
    if selected_note:
        encounter_id = int(selected_note.split(' | ')[0])
        note_row = notes_df[notes_df['id'] == encounter_id].iloc[0]
        st.markdown('#### Nota clínica')
        note_text = note_row.get('note_text') or ''
        if note_text.strip():
            st.text_area('Contenido de la nota', value=note_text, height=360, disabled=True, key=f'patient_note_text_{encounter_id}')
        else:
            st.info('Esta consulta no tiene nota clínica guardada.')
        if note_row.get('note_json'):
            try:
                parsed_note = json.loads(note_row.get('note_json') or '{}')
                with st.expander('Resumen estructurado'):
                    st.json(parsed_note)
            except Exception:
                pass


def render_patient_list_page(user:dict):
    render_page_header('Historial atenciones de pacientes', 'Busca pacientes, abre su ficha y revisa su historial de atenciones en el sistema.')
    st.info('Selecciona un paciente de la tabla para crear una nueva atención o revisar su ficha e historial.')
    if st.button('Crear paciente nuevo'):
        clear_patient_form_state()
        st.session_state['active_page'] = 'Atención clínica'
        st.rerun()

    search_term = st.text_input('Buscar paciente por MRN, RUN, nombre, teléfono, correo, dirección, tutor, centro o fecha de nacimiento', value='')
    patients = search_patients(search_term, limit=200) if search_term.strip() else get_patients_df()
    if patients.empty:
        st.warning('No hay pacientes registrados.')
    else:
        patients = patients.assign(paciente=patients['first_name'].fillna('') + ' ' + patients['last_name'].fillna(''))
        display = patients[['id','centro','mrn','national_id','paciente','birth_date','sex','updated_at']].rename(columns={
            'centro':'Centro','national_id':'RUN','paciente':'Paciente','birth_date':'Nacimiento','updated_at':'Última actualización'
        })
        patient_action_display = display.copy()
        patient_action_display.insert(0, 'Seleccionar', False)
        edited_patients = st.data_editor(
            patient_action_display,
            use_container_width=True,
            hide_index=True,
            disabled=[c for c in patient_action_display.columns if c != 'Seleccionar'],
            column_config={
                'Seleccionar': st.column_config.CheckboxColumn('Seleccionar', help='Selecciona un paciente para crear una atención o revisar su historial')
            },
            key='patients_action_table',
        )
        selected_patient_rows = edited_patients.loc[edited_patients['Seleccionar']]
        if len(selected_patient_rows) > 1:
            st.warning('Selecciona solo un paciente para continuar.')
        selected_patient_id = int(selected_patient_rows.iloc[0]['id']) if len(selected_patient_rows) == 1 else st.session_state.get('patient_id')
        if selected_patient_id:
            patient_id = selected_patient_id
            patient = get_patient(patient_id)
            if patient:
                st.session_state['patient_id'] = patient_id
                st.session_state['patient_facility_id'] = patient.get('facility_id')
                with st.container(border=True):
                    st.markdown('### Paciente seleccionado')
                    d1,d2,d3=st.columns(3)
                    with d1:
                        st.write(f"**MRN:** {patient.get('mrn')}")
                        st.write(f"**Nombre:** {patient.get('first_name','')} {patient.get('last_name','')}")
                    with d2:
                        st.write(f"**Centro:** {patient.get('facility_name') or 'Sin centro'}")
                        st.write(f"**RUN:** {patient.get('national_id','')}")
                    with d3:
                        st.write(f"**Teléfono:** {patient.get('phone','')}")
                        st.write(f"**Tutor/Responsable:** {patient.get('guardian_name','')}")
                    if patient.get('address'):
                        st.write(f"**Dirección:** {patient.get('address','')}")
                    action_new, action_history = st.columns(2)
                    with action_new:
                        if st.button('Crear nueva atención', key='patient_create_new_attention'):
                            st.session_state['patient_id'] = patient_id
                            st.session_state['patient_facility_id'] = patient.get('facility_id') if patient else None
                            encounter_id = create_encounter(patient_id, {
                                "visit_type": "INGRESO",
                                "brought_by": "",
                                "informant": "",
                                "provider_name": user.get('full_name') or user.get('email') or "",
                                "facility_name": "",
                                "chief_complaint": "",
                            })
                            st.session_state['encounter_id'] = encounter_id
                            st.session_state['active_page'] = 'Atención clínica'
                            st.rerun()
                    with action_history:
                        if st.button('Ver ficha e historial', key='patient_open_chart_history'):
                            st.session_state['patient_id'] = patient_id
                            st.session_state['patient_facility_id'] = patient.get('facility_id')
                            st.session_state['active_page'] = 'Pacientes'
                            st.rerun()
                render_patient_history_detail(patient_id)
        else:
            st.caption('Selecciona un paciente en la tabla para habilitar acciones.')

def render_patient_chart_page(user:dict):
    render_patient_list_page(user)
    return
    patient_id = st.session_state.get('patient_id')
    render_page_header('Ficha del paciente')
    if not patient_id:
        st.info('No hay paciente seleccionado.')
        return
    patient = get_patient(patient_id)
    if not patient:
        st.warning('Paciente no encontrado.')
        return

    if st.button('Volver a Agenda'):
        st.session_state['active_page'] = 'Agenda'
        st.rerun()

    st.markdown('### Datos del paciente')
    d1,d2,d3=st.columns(3)
    with d1:
        st.write(f"**ID/MRN:** {patient.get('mrn') or ''}")
        st.write(f"**Nombre:** {patient.get('first_name','')} {patient.get('last_name','')}")
        st.write(f"**Centro:** {patient.get('facility_name') or 'Sin centro'}")
    with d2:
        st.write(f"**RUN/documento:** {patient.get('national_id') or ''}")
        st.write(f"**Nacimiento:** {patient.get('birth_date') or ''}")
        st.write(f"**Sexo:** {patient.get('sex') or ''}")
    with d3:
        st.write(f"**Teléfono:** {patient.get('phone') or ''}")
        st.write(f"**Email:** {patient.get('email') or ''}")
        st.write(f"**Tutor/responsable:** {patient.get('guardian_name') or ''}")
    st.write(f"**Dirección:** {patient.get('address') or ''}")

    notes_df = get_patient_notes_df(patient_id)
    st.markdown('### Resúmenes y notas creadas')
    if notes_df.empty:
        st.info('No hay resúmenes ni notas creadas para este paciente.')
        return

    summary_df = notes_df[['id','created_at','visit_type','chief_complaint','status','updated_at']].rename(columns={
        'id':'ID','created_at':'Fecha','visit_type':'Tipo','chief_complaint':'Motivo','status':'Estado','updated_at':'Actualizado'
    })
    st.dataframe(summary_df, use_container_width=True, hide_index=True)

    note_options = notes_df.apply(lambda r: f"{r['id']} | {r['created_at']} | {r.get('chief_complaint') or 'Sin motivo'}", axis=1).tolist()
    selected_note = st.selectbox('Seleccionar nota para visualizar', [''] + note_options)
    if selected_note:
        encounter_id = int(selected_note.split(' | ')[0])
        note_row = notes_df[notes_df['id'] == encounter_id].iloc[0]
        st.markdown('#### Nota clínica')
        note_text = note_row.get('note_text') or ''
        if note_text.strip():
            st.text_area('Contenido de la nota', value=note_text, height=360, disabled=True)
        else:
            st.info('Esta consulta no tiene nota clínica guardada.')
        if note_row.get('note_json'):
            try:
                parsed_note = json.loads(note_row.get('note_json') or '{}')
                with st.expander('Resumen estructurado'):
                    st.json(parsed_note)
            except Exception:
                pass


def render_admin_page(user:dict):
    if not is_admin_user(user): st.error('Solo el perfil administrador puede acceder.'); st.stop()
    render_page_header('Administrador de la herramienta', 'Gestiona usuarios, centros, agendas y configuración operativa.')
    tab1,tab2,tab3=st.tabs(['Usuarios y médicos','Centros','Configuración correo'])
    with tab1:
        st.markdown('#### Crear nuevo profesional o administrador')
        with st.form('new_user_form'):
            email=st.text_input('Correo del usuario'); full_name=st.text_input('Nombre completo'); phone=st.text_input('Teléfono'); specialty=st.text_input('Especialidad / rol'); role=st.selectbox('Perfil',[DOCTOR_ROLE,ADMIN_ROLE]); professional_id=st.text_input('Registro profesional / ID interno'); password=st.text_input('Contraseña temporal', type='password'); submitted=st.form_submit_button('Crear usuario')
        if submitted:
            if not email.strip() or not password.strip():
                st.error('Correo y contraseña temporal son obligatorios.')
            else:
                try:
                    create_user(email,password,role,full_name,phone,specialty)
                    update_user(get_user_by_email(email)['id'], email, full_name, phone, specialty, role, 1, professional_id, 'SOAP')
                    st.success('Usuario creado. En su primer ingreso deberá completar/confirmar perfil.')
                except Exception as e:
                    st.error(f'No se pudo crear usuario: {e}')

        st.markdown('#### Lista de usuarios y edición rápida')
        users = get_users_df()
        if users.empty:
            st.info('No hay usuarios registrados.')
        else:
            st.dataframe(users.rename(columns={
                'id':'ID','email':'Correo','role':'Perfil','full_name':'Nombre','phone':'Teléfono','specialty':'Especialidad','active':'Activo','updated_at':'Actualizado','professional_id':'Registro profesional','preferred_note_format':'Formato nota'}), use_container_width=True, hide_index=True)
            selected_user = st.selectbox('Seleccionar usuario para editar', [''] + users.apply(lambda r: f"{r['id']} | {r['full_name'] or r['email']} ({r['role']})", axis=1).tolist(), key='edit_user_select')
            if selected_user:
                user_id = int(selected_user.split(' | ')[0])
                user_to_edit = get_user(user_id)
                if user_to_edit:
                    with st.form('edit_user_form'):
                        edit_email = st.text_input('Correo destino del profesional', value=user_to_edit.get('email') or '', help='Este correo se usará como destino cuando este usuario exporte y envíe una nota clínica.')
                        edit_full_name = st.text_input('Nombre completo', value=user_to_edit.get('full_name') or '')
                        edit_phone = st.text_input('Teléfono', value=user_to_edit.get('phone') or '')
                        edit_specialty = st.text_input('Especialidad / rol', value=user_to_edit.get('specialty') or '')
                        edit_role = st.selectbox('Perfil', [DOCTOR_ROLE,ADMIN_ROLE], index=0 if user_to_edit.get('role') != ADMIN_ROLE else 1)
                        edit_active = st.checkbox('Activo', value=bool(user_to_edit.get('active',1)))
                        edit_professional_id = st.text_input('Registro profesional / ID interno', value=user_to_edit.get('professional_id') or '')
                        edit_preferred_format = st.selectbox('Formato nota preferido', ['SOAP','Ficha clínica completa'], index=0 if user_to_edit.get('preferred_note_format','SOAP') == 'SOAP' else 1)
                        save_user = st.form_submit_button('Guardar cambios de usuario')
                    if save_user:
                        try:
                            update_user(user_id, edit_email, edit_full_name, edit_phone, edit_specialty, edit_role, 1 if edit_active else 0, edit_professional_id, edit_preferred_format)
                            if st.session_state.get('auth_user', {}).get('id') == user_id:
                                st.session_state['auth_user'] = get_user(user_id)
                            st.success('Usuario actualizado correctamente.')
                            st.rerun()
                        except Exception as e:
                            st.error(f'No se pudo actualizar el usuario: {e}')
        providers = users[users['role'].isin([DOCTOR_ROLE,ADMIN_ROLE])] if not users.empty else users
        if not providers.empty:
            provider_select = st.selectbox('Seleccionar profesional', [''] + providers.apply(lambda r: f"{r['id']} | {r['full_name'] or r['email']}", axis=1).tolist(), key='provider_history_select')
            if provider_select:
                provider_id = int(provider_select.split(' | ')[0])
                st.markdown('##### Agenda asignada')
                provider_appts = get_appointments_df('', None, provider_id)
                if provider_appts.empty:
                    st.info('No hay citas asignadas a este profesional.')
                else:
                    st.dataframe(provider_appts[['id','fecha','hora','centro','paciente','estado','motivo']].rename(columns={'id':'ID','fecha':'Fecha','hora':'Hora','centro':'Centro','paciente':'Paciente','estado':'Estado','motivo':'Motivo'}), use_container_width=True, hide_index=True)
                st.markdown('##### Encuentros clínicos')
                provider_enc = get_provider_encounters_df(provider_id)
                if provider_enc.empty:
                    st.info('No hay encuentros clínicos asignados a este profesional.')
                else:
                    st.dataframe(provider_enc[['id','mrn','patient_name','visit_type','chief_complaint','status','created_at']].rename(columns={'id':'ID','mrn':'MRN','patient_name':'Paciente','visit_type':'Tipo','chief_complaint':'Motivo','status':'Estado','created_at':'Creado'}), use_container_width=True, hide_index=True)
        else:
            st.info('No hay profesionales disponibles para historial.')
    with tab2:
        with st.form('facility_form'):
            name=st.text_input('Nombre centro'); code=st.text_input('Sigla centro', help='Ejemplo: Centro Jorge 1 = CJ1. Se usará para generar IDs como CJ100001.'); address=st.text_input('Dirección'); phone=st.text_input('Teléfono'); ok=st.form_submit_button('Crear centro')
        if ok:
            if not name.strip() or not code.strip():
                st.error('El nombre y la sigla del centro son obligatorios.')
            else:
                try:
                    create_facility(name,code,address,phone); st.success('Centro creado.')
                except Exception as e:
                    st.error(f'No se pudo crear el centro: {e}')
        facilities_admin = get_facilities_df(active_only=False)
        st.dataframe(facilities_admin, use_container_width=True, hide_index=True)
        if not facilities_admin.empty:
            selected_facility = st.selectbox('Seleccionar centro para editar', [''] + facilities_admin.apply(lambda r: f"{r['id']} | {r['name']} ({r.get('code') or 'sin sigla'})", axis=1).tolist(), key='edit_facility_select')
            if selected_facility:
                facility_id_to_edit = int(selected_facility.split(' | ')[0])
                facility_row = facilities_admin[facilities_admin['id'] == facility_id_to_edit].iloc[0]
                with st.form('edit_facility_form'):
                    edit_facility_name = st.text_input('Nombre centro', value=facility_row.get('name') or '')
                    edit_facility_code = st.text_input('Sigla centro', value=facility_row.get('code') or '')
                    edit_facility_address = st.text_input('Dirección', value=facility_row.get('address') or '')
                    edit_facility_phone = st.text_input('Teléfono', value=facility_row.get('phone') or '')
                    edit_facility_active = st.checkbox('Activo', value=bool(facility_row.get('active', 1)))
                    save_facility = st.form_submit_button('Guardar cambios de centro')
                if save_facility:
                    try:
                        update_facility(facility_id_to_edit, edit_facility_name, edit_facility_code, edit_facility_address, edit_facility_phone, 1 if edit_facility_active else 0)
                        st.success('Centro actualizado correctamente.')
                        st.rerun()
                    except Exception as e:
                        st.error(f'No se pudo actualizar el centro: {e}')
    with tab3:
        cfg = get_smtp_config()
        auto_email = get_bool_setting("AUTO_EMAIL_ON_NOTE", False)
        st.code(
            f"SMTP_ENABLED={str(cfg['enabled']).lower()}\n"
            f"SMTP_HOST={cfg['host']}\n"
            f"SMTP_PORT={cfg['port']}\n"
            f"SMTP_USER={cfg['user']}\n"
            f"SMTP_PASSWORD={mask_secret(cfg['password'])}\n"
            f"SMTP_FROM={cfg['from']}\n"
            f"AUTO_EMAIL_ON_NOTE={str(auto_email).lower()}\n\n"
            "Destino: correo del profesional autenticado que presiona el botón"
        )
        st.info('Correo configurado.' if cfg["user"] and cfg["password"] else 'Correo aún no configurado.')

# =========================================================
# Inicialización
# =========================================================
validate_startup_security()
init_db()
seed_cie10_if_empty()
try:
    ensure_default_admin_and_facility()
except RuntimeError as exc:
    st.error(str(exc))
    st.stop()

def clear_patient_form_state():
    st.session_state['patient_id'] = None
    st.session_state['patient_facility_id'] = None
    st.session_state['patient_mrn'] = ""
    st.session_state['patient_first_name'] = ""
    st.session_state['patient_last_name'] = ""
    st.session_state['patient_national_id'] = ""
    st.session_state['patient_birth_date'] = ""
    st.session_state['patient_sex'] = ""
    st.session_state['patient_phone'] = ""
    st.session_state['patient_email'] = ""
    st.session_state['patient_guardian_name'] = ""
    st.session_state['patient_address'] = ""

for key, default in [
    ("patient_id", None), ("encounter_id", None), ("appointment_id", None), ("active_appointment_id", None), ("transcript_text", ""), ("segments", []),
    ("patient_facility_id", None),
    ("dictation_text", ""), ("dictation_segments", []), ("images_text", ""),
    ("note_json", None), ("note_txt", ""), ("organized_text", ""), ("pending_note_send", None), ("pending_same_day_schedule", None),
    ("patient_search", ""), ("patient_search_select", ""),
    ("patient_mrn", ""), ("patient_first_name", ""), ("patient_last_name", ""),
    ("patient_national_id", ""), ("patient_birth_date", ""), ("patient_sex", ""),
    ("patient_phone", ""), ("patient_email", ""), ("patient_guardian_name", ""),
    ("patient_address", ""), ("return_to", None), ("agenda_message", None), ("pending_appointment", None),
]:
    st.session_state.setdefault(key, default)

# =========================================================
# UI
# =========================================================
auth_user = require_login()
require_profile_completion(auth_user)

if st.session_state.get('active_page') is None:
    render_main_menu(auth_user)
page = st.session_state.get('active_page', 'Agenda')

with st.sidebar:
    render_app_brand(compact=True)
    render_info_card('Usuario', f"{auth_user.get('full_name') or auth_user.get('email')} | Perfil: {auth_user.get('role')}", '#0F4C81')

    st.markdown('---')
    st.subheader('Navegacion')
    render_sidebar_nav_button('Menú principal', None, 'nav_menu_principal')
    render_sidebar_nav_button('Agenda', 'Agenda', 'nav_agenda')
    render_sidebar_nav_button('Atención clínica', 'Atención clínica', 'nav_atencion_clinica')
    render_sidebar_nav_button('Pacientes e historial', 'Pacientes', 'nav_pacientes')
    if is_admin_user(auth_user):
        render_sidebar_nav_button('Administrador', 'Administrador', 'nav_administrador')
    current_active = get_active_page_label()
    st.markdown(
        f'<span class="as-current-page-pill">Actual: {html.escape(str(current_active))}</span>',
        unsafe_allow_html=True,
    )

    st.markdown('---')
    st.subheader("Contexto")
    active_patient = get_patient(st.session_state.get("patient_id")) if st.session_state.get("patient_id") else None
    active_encounter = get_encounter(st.session_state.get("encounter_id")) if st.session_state.get("encounter_id") else None
    if active_patient:
        st.write(f"Paciente: {active_patient.get('mrn') or active_patient.get('id')} - {active_patient.get('first_name','')} {active_patient.get('last_name','')}")
    else:
        st.caption("Sin paciente activo.")
    if active_encounter:
        status_label = "Revisado" if is_physician_reviewed(active_encounter) else "Borrador IA"
        st.write(f"Consulta: #{active_encounter.get('id')} | {status_label}")
    else:
        st.caption("Sin consulta activa.")

    st.markdown('---')
    st.subheader("Integraciones")
    hl7_sidebar_settings = HL7Settings.from_env()
    fhir_sidebar_enabled = get_bool_setting("FHIR_ENABLED", False)
    st.caption(f"HL7: {hl7_sidebar_settings.transport_label}")
    st.caption(f"FHIR: {'activado' if fhir_sidebar_enabled else 'solo export local'}")
    if google_login_enabled():
        st.caption(f"Google OAuth: {'configurado' if google_oauth_configured() else 'pendiente'}")

    st.markdown('---')
    if st.button('Cerrar sesión', key='logout'):
        st.session_state.clear()
        st.rerun()

if page == 'Agenda':
    render_agenda_page(auth_user); st.stop()
if page == 'Pacientes':
    render_patient_list_page(auth_user); st.stop()
if page == 'Ficha paciente':
    render_patient_chart_page(auth_user); st.stop()
if page == 'Administrador':
    render_admin_page(auth_user); st.stop()

render_page_header('Atención clínica ambulatoria', 'Ficha clínica local: pacientes, CIE-10, nota IA, exportación clínica, correo automático y HL7 v2 preparado para integración HIS.')

# -------------------------
# 1. Pacientes
# -------------------------
if st.session_state.get("appointment_id"):
    appt_id = st.session_state.pop("appointment_id")
    st.session_state["active_appointment_id"] = appt_id
    if st.session_state.get("patient_id"):
        patient = get_patient(st.session_state["patient_id"])
        if patient:
            st.session_state["patient_mrn"] = patient.get("mrn") or ""
            st.session_state["patient_facility_id"] = patient.get("facility_id")
            st.session_state["patient_first_name"] = patient.get("first_name") or ""
            st.session_state["patient_last_name"] = patient.get("last_name") or ""
            st.session_state["patient_national_id"] = patient.get("national_id") or ""
            st.session_state["patient_birth_date"] = patient.get("birth_date") or ""
            st.session_state["patient_sex"] = patient.get("sex") or ""
            st.session_state["patient_phone"] = patient.get("phone") or ""
            st.session_state["patient_email"] = patient.get("email") or ""
            st.session_state["patient_guardian_name"] = patient.get("guardian_name") or ""
            st.session_state["patient_address"] = patient.get("address") or ""

with st.expander("1) Paciente y búsqueda rápida", expanded=True):
    st.caption("Busca por MRN, RUN o nombre para detectar si el paciente ya existe. Si existe, cargaremos su ficha y su historial clínico.")
    search_term = st.text_input("Buscar paciente (MRN / RUN / nombre)", value=st.session_state.get("patient_search", ""), key="patient_search")
    matches = search_patients(search_term, limit=20) if search_term.strip() else pd.DataFrame()
    if not matches.empty:
        options = matches.apply(lambda r: f"{r['id']} | {r['mrn']} | {r['national_id']} | {r['first_name']} {r['last_name']} | {r.get('centro') or 'Sin centro'} ({r['birth_date'] or 'sin fecha'})", axis=1).tolist()
        selected_match = st.selectbox("Coincidencias encontradas", [""] + options, key="patient_search_select")
        if selected_match:
            selected_id = int(selected_match.split(" | ")[0])
            patient = get_patient(selected_id)
            if patient:
                st.session_state["patient_id"] = selected_id
                st.session_state["patient_facility_id"] = patient.get("facility_id")
                st.session_state["patient_mrn"] = patient.get("mrn") or ""
                st.session_state["patient_first_name"] = patient.get("first_name") or ""
                st.session_state["patient_last_name"] = patient.get("last_name") or ""
                st.session_state["patient_national_id"] = patient.get("national_id") or ""
                st.session_state["patient_birth_date"] = patient.get("birth_date") or ""
                st.session_state["patient_sex"] = patient.get("sex") or ""
                st.session_state["patient_phone"] = patient.get("phone") or ""
                st.session_state["patient_email"] = patient.get("email") or ""
                st.session_state["patient_guardian_name"] = patient.get("guardian_name") or ""
                st.session_state["patient_address"] = patient.get("address") or ""
                st.success("Paciente existente cargado. Revisa los datos y actualiza si es necesario.")

    if st.session_state.get("patient_id"):
        patient = get_patient(st.session_state["patient_id"])
        if patient:
            st.info(f"Paciente activo: {patient.get('mrn')} - {patient.get('first_name', '')} {patient.get('last_name', '')}")
            if patient.get("facility_name"):
                st.caption(f"Centro del paciente: {patient.get('facility_name')}")
            enc_df = get_encounters_df(patient.get("id"))
            if not enc_df.empty:
                st.markdown("**Historial clínico reciente**")
                st.dataframe(enc_df[['id','created_at','visit_type','chief_complaint','status']], use_container_width=True, hide_index=True)
            else:
                st.info("Aún no hay encuentros guardados para este paciente.")
    else:
        st.info("Paciente nuevo. Completa los datos y guarda la ficha para crear un registro.")

    facilities_for_patient = get_facilities_df()
    if facilities_for_patient.empty:
        st.error("Primero debes crear un centro en Administrador > Centros.")
        st.stop()
    facility_options = facilities_for_patient.apply(lambda r: f"{r['id']} - {r['name']}", axis=1).tolist()
    current_facility_id = st.session_state.get("patient_facility_id")
    default_facility_index = 0
    if current_facility_id:
        matches_facility = [i for i, label in enumerate(facility_options) if label.startswith(f"{int(current_facility_id)} - ")]
        default_facility_index = matches_facility[0] if matches_facility else 0
    selected_patient_facility = st.selectbox("Centro al que pertenece el paciente", facility_options, index=default_facility_index)
    selected_patient_facility_id = int(selected_patient_facility.split(" - ")[0])
    st.session_state["patient_facility_id"] = selected_patient_facility_id

    c1, c2, c3 = st.columns(3)
    with c1:
        current_patient_id = st.session_state.get("patient_id")
        mrn_preview = st.session_state.get("patient_mrn", "") if current_patient_id else preview_next_patient_mrn(selected_patient_facility_id)
        st.text_input("Identificador automático / MRN", value=mrn_preview, disabled=True)
        mrn = st.session_state.get("patient_mrn", "") if current_patient_id else None
        first_name = st.text_input("Nombre", value=st.session_state.get("patient_first_name", ""), key="patient_first_name")
        last_name = st.text_input("Apellidos", value=st.session_state.get("patient_last_name", ""), key="patient_last_name")
    with c2:
        national_id = st.text_input("RUN / documento", value=st.session_state.get("patient_national_id", ""), key="patient_national_id")
        birth_date_state = st.session_state.get("patient_birth_date", "")
        if isinstance(birth_date_state, date):
            birth_date_value = birth_date_state
        else:
            try:
                birth_date_value = date.fromisoformat(birth_date_state) if birth_date_state else None
            except Exception:
                birth_date_value = None
        birth_date = st.date_input("Fecha de nacimiento", value=birth_date_value, min_value=date(1900, 1, 1), max_value=date.today(), format="DD/MM/YYYY", key="patient_birth_date")
        sex_options = ["", "M", "F", "O", "U"]
        current_sex = st.session_state.get("patient_sex", "")
        sex = st.selectbox("Sexo", sex_options, index=sex_options.index(current_sex) if current_sex in sex_options else 0, help="M masculino, F femenino, O otro, U desconocido/no informado", key="patient_sex")
    with c3:
        phone = st.text_input("Teléfono", value=st.session_state.get("patient_phone", ""), key="patient_phone")
        email = st.text_input("Email", value=st.session_state.get("patient_email", ""), key="patient_email")
        guardian_name = st.text_input("Adulto responsable / tutor", value=st.session_state.get("patient_guardian_name", ""), key="patient_guardian_name")
    address = st.text_input("Dirección", value=st.session_state.get("patient_address", ""), key="patient_address")

    schedule_on_patient_save = st.checkbox("Agendar al guardar en un cupo disponible", value=False, key="clinical_schedule_on_save")
    selected_clinical_slot = None
    clinical_schedule_reason = ""
    clinical_start_attention = False
    if schedule_on_patient_save:
        st.markdown("**Cupo de agenda para este paciente**")
        providers_for_schedule = get_users_df()
        if not providers_for_schedule.empty:
            providers_for_schedule = providers_for_schedule[
                (providers_for_schedule["role"].isin([DOCTOR_ROLE, ADMIN_ROLE])) &
                (providers_for_schedule["active"] == 1)
            ].copy()
        if providers_for_schedule.empty:
            st.info("No hay profesionales activos para agendar.")
        else:
            sched_c1, sched_c2, sched_c3 = st.columns(3)
            with sched_c1:
                clinical_schedule_day = st.date_input("Dia del cupo", value=date.today(), min_value=date(1970, 1, 1), format="DD/MM/YYYY", key="clinical_schedule_day")
            with sched_c2:
                if auth_user.get("role") == DOCTOR_ROLE:
                    selected_schedule_provider_id = int(auth_user["id"])
                    st.text_input("Profesional", value=auth_user.get("full_name") or auth_user.get("email") or "", disabled=True, key="clinical_schedule_provider_locked")
                else:
                    provider_labels = []
                    for _, provider_row in providers_for_schedule.iterrows():
                        provider_name = provider_row.get("full_name")
                        if pd.isna(provider_name) or not str(provider_name).strip():
                            provider_name = provider_row.get("email") or ""
                        provider_labels.append(f"{int(provider_row['id'])} - {provider_name}")
                    provider_default = 0
                    current_provider_matches = [i for i, label in enumerate(provider_labels) if label.startswith(f"{int(auth_user.get('id') or 0)} - ")]
                    if current_provider_matches:
                        provider_default = current_provider_matches[0]
                    selected_provider_label = st.selectbox("Profesional", provider_labels, index=provider_default, key="clinical_schedule_provider")
                    selected_schedule_provider_id = int(selected_provider_label.split(" - ")[0])
            with sched_c3:
                clinical_slot_minutes = st.selectbox("Duracion cupo", [15, 20, 30, 45, 60], index=2, key="clinical_schedule_slot_minutes")

            time_c1, time_c2 = st.columns(2)
            with time_c1:
                clinical_agenda_start = st.time_input("Inicio jornada", value=dt_time(8, 0), key="clinical_schedule_start")
            with time_c2:
                clinical_agenda_end = st.time_input("Fin jornada", value=dt_time(18, 0), key="clinical_schedule_end")
            clinical_schedule_reason = st.text_input("Motivo / observacion de la cita", value="", key="clinical_schedule_reason")
            clinical_start_attention = st.checkbox("Crear registro clinico e iniciar atencion al guardar", value=True, key="clinical_schedule_start_attention")

            selected_facility_df = facilities_for_patient[facilities_for_patient["id"] == selected_patient_facility_id]
            selected_provider_df = providers_for_schedule[providers_for_schedule["id"] == selected_schedule_provider_id]
            appts_for_slots = get_appointments_df(clinical_schedule_day.isoformat(), selected_patient_facility_id, selected_schedule_provider_id)
            clinical_slots = build_daily_slots_df(
                clinical_schedule_day,
                selected_facility_df,
                selected_provider_df,
                appts_for_slots,
                clinical_agenda_start.strftime("%H:%M"),
                clinical_agenda_end.strftime("%H:%M"),
                int(clinical_slot_minutes),
            )
            available_clinical_slots = clinical_slots[clinical_slots["Cupo"] == "Disponible"].reset_index(drop=True) if not clinical_slots.empty else pd.DataFrame()
            if available_clinical_slots.empty:
                st.warning("No hay cupos disponibles con estos filtros.")
            else:
                visible_clinical_slots = available_clinical_slots[["Centro", "Profesional", "Hora", "Cupo"]]
                clinical_slot_selection = st.dataframe(
                    visible_clinical_slots,
                    use_container_width=True,
                    hide_index=True,
                    on_select="rerun",
                    selection_mode="single-row",
                    key="clinical_available_slot_table",
                )
                clinical_selected_rows = clinical_slot_selection.selection.rows if clinical_slot_selection and clinical_slot_selection.selection else []
                if clinical_selected_rows:
                    selected_row_idx = int(clinical_selected_rows[0])
                    if selected_row_idx < len(available_clinical_slots):
                        selected_clinical_slot = available_clinical_slots.iloc[selected_row_idx]
                    else:
                        st.info("La lista de cupos cambio. Selecciona nuevamente un cupo disponible.")

    save_patient_col, history_patient_col = st.columns([1, 1])
    with save_patient_col:
        save_patient_clicked = st.button("Guardar / actualizar paciente")
    with history_patient_col:
        if st.button("Ver historial de atenciones", disabled=not st.session_state.get("patient_id")):
            st.session_state["active_page"] = "Pacientes"
            st.rerun()

    if save_patient_clicked:
        if schedule_on_patient_save and selected_clinical_slot is None:
            st.error("Selecciona un cupo disponible antes de guardar con agendamiento.")
            st.stop()
        try:
            patient_id = create_or_update_patient({
                "id": st.session_state.get("patient_id"),
                "mrn": (mrn or "").strip() or None,
                "first_name": first_name,
                "last_name": last_name,
                "national_id": national_id,
                "birth_date": birth_date.isoformat() if birth_date else "",
                "sex": sex,
                "phone": phone,
                "email": email,
                "address": address,
                "guardian_name": guardian_name,
                "facility_id": selected_patient_facility_id,
            })
            st.session_state["patient_id"] = patient_id
            st.session_state["patient_facility_id"] = selected_patient_facility_id
            saved_patient = get_patient(patient_id)
            if saved_patient:
                st.session_state["patient_mrn"] = saved_patient.get("mrn") or ""
            st.success(f"Paciente activo guardado. Identificador: {saved_patient.get('mrn') if saved_patient else patient_id}")
            if schedule_on_patient_save and selected_clinical_slot is not None:
                appointment_id = create_appointment(
                    patient_id,
                    int(selected_clinical_slot["_facility_id"]),
                    int(selected_clinical_slot["_provider_user_id"]),
                    selected_clinical_slot["_fecha"],
                    selected_clinical_slot["Hora"],
                    int(selected_clinical_slot["_duration_minutes"]),
                    clinical_schedule_reason,
                )
                st.session_state["active_appointment_id"] = appointment_id
                st.session_state["agenda_message"] = f"Cita agendada para {selected_clinical_slot['_fecha']} {selected_clinical_slot['Hora']}."
                st.session_state["pending_appointment"] = None
                st.session_state["return_to"] = None
                if clinical_start_attention:
                    encounter_id = create_encounter(patient_id, {
                        "visit_type": "INGRESO",
                        "brought_by": "",
                        "informant": "",
                        "provider_name": selected_clinical_slot.get("Profesional") or auth_user.get("full_name") or auth_user.get("email") or "",
                        "facility_name": selected_clinical_slot.get("Centro") or "",
                        "chief_complaint": clinical_schedule_reason,
                    })
                    attach_encounter_to_appointment(appointment_id, encounter_id)
                    set_appointment_status(appointment_id, "en_atencion")
                    st.session_state["encounter_id"] = encounter_id
                    st.success(f"Cita agendada y registro clinico creado. Consulta activa: #{encounter_id}")
                else:
                    st.success(f"Cita agendada para {selected_clinical_slot['_fecha']} {selected_clinical_slot['Hora']}.")
            pending_appt = st.session_state.get("pending_appointment")
            if pending_appt:
                create_appointment(
                    patient_id,
                    pending_appt["facility_id"],
                    pending_appt["provider_user_id"],
                    pending_appt["appointment_date"],
                    pending_appt["appointment_time"],
                    int(pending_appt["duration_minutes"]),
                    pending_appt.get("reason", ""),
                )
                st.session_state["pending_appointment"] = None
                st.session_state["return_to"] = None
                st.session_state["agenda_message"] = f"Paciente creado y cita agendada para {pending_appt['appointment_date']} {pending_appt['appointment_time']}."
                st.session_state["active_page"] = "Agenda"
                st.rerun()
            if st.session_state.get("return_to"):
                next_page = st.session_state.pop("return_to")
                st.session_state["agenda_message"] = f"Paciente creado y vuelto a Agenda. ID: {patient_id}"
                st.session_state["active_page"] = next_page
                st.rerun()
        except Exception as e:
            st.error(f"No se pudo guardar el paciente: {e}")

# -------------------------
# 2. Encuentro ambulatorio
# -------------------------
with st.expander("2) Consulta ambulatoria", expanded=True):
    if not st.session_state.get("patient_id"):
        st.warning("Primero guarda o selecciona un paciente.")
    else:
        ec1, ec2, ec3, ec4 = st.columns(4)
        with ec1:
            visit_type = st.selectbox("Tipo de consulta", ["INGRESO", "CONTROL", "OTRO"])
        with ec2:
            brought_by = st.selectbox("Traído por", ["MADRE", "PADRE", "AMBOS PADRES", "OTRO", "NO APLICA"])
        with ec3:
            provider_name = st.text_input("Profesional", value=auth_user.get("full_name") or "Dr/a. Jorge")
        with ec4:
            facility_name = st.text_input("Centro", value="Consulta ambulatoria")
        informant = st.text_input("Informante", value="")
        chief_complaint = st.text_area("Motivo de consulta inicial", height=90)

        col_btn, col_spacer = st.columns([1, 4])
        with col_btn:
            if st.button("Nuevo registro", help="Nueva atención dentro de este registro"):
                encounter_id = create_encounter(st.session_state["patient_id"], {
                    "visit_type": visit_type,
                    "brought_by": brought_by,
                    "informant": informant,
                    "provider_name": provider_name,
                    "facility_name": facility_name,
                    "chief_complaint": chief_complaint,
                })
                st.session_state["encounter_id"] = encounter_id
                st.success(f"Registro clínico creado. ID: {encounter_id}")

        enc_df = get_encounters_df(st.session_state.get("patient_id"))
        if not enc_df.empty:
            enc_df["label"] = enc_df.apply(lambda r: f"#{r['id']} - {r['created_at']} - {r['visit_type']} - {r['chief_complaint'] or ''}"[:140], axis=1)
            selected_enc = st.selectbox("O seleccionar consulta existente", [""] + enc_df["label"].tolist())
            if selected_enc:
                st.session_state["encounter_id"] = int(enc_df.loc[enc_df["label"] == selected_enc, "id"].iloc[0])

        if st.session_state.get("encounter_id"):
            st.info(f"Consulta activa: #{st.session_state['encounter_id']}")

# -------------------------
# 3. Audio e IA
# -------------------------
render_info_card("Funciones clínicas principales", "Graba, transcribe, estructura, revisa y exporta la atención activa.", "#1B998B")
col1, col2 = st.columns(2)
with col1:
    with st.expander("3.1 Grabar conversación", expanded=True):
        minutes_conv = st.slider("Minutos conversación", 1, 30, 5, key="minutes_conv")
        if st.button("Grabar conversación"):
            if not st.session_state.get("patient_id"):
                st.error("Primero guarda o selecciona un paciente para asociar el audio.")
            elif not SOUNDDEVICE_AVAILABLE:
                st.warning("La grabacion local no esta disponible porque sounddevice/PortAudio no se pudo cargar. En Streamlit Cloud sube un audio en la seccion 3.3.")
            else:
                wav_path = patient_audio_path(st.session_state.get("patient_id"), "conversacion")
                with st.spinner("Grabando conversación..."):
                    record_wav(wav_path, seconds=int(minutes_conv * 60))
                st.session_state["wav_conv"] = wav_path
                st.success(f"Guardado: {audio_file_label(wav_path)}")

with col2:
    with st.expander("3.2 Grabar dictado del médico", expanded=True):
        minutes_dict = st.slider("Minutos dictado", 1, 30, 2, key="minutes_dict")
        if st.button("Grabar dictado"):
            if not st.session_state.get("patient_id"):
                st.error("Primero guarda o selecciona un paciente para asociar el audio.")
            elif not SOUNDDEVICE_AVAILABLE:
                st.warning("La grabacion local no esta disponible porque sounddevice/PortAudio no se pudo cargar. En Streamlit Cloud sube un audio en la seccion 3.3.")
            else:
                wav_path = patient_audio_path(st.session_state.get("patient_id"), "dictado")
                with st.spinner("Grabando dictado..."):
                    record_wav(wav_path, seconds=int(minutes_dict * 60))
                st.session_state["wav_dict"] = wav_path
                st.success(f"Guardado: {audio_file_label(wav_path)}")

with st.expander("Herramientas de audio Windows", expanded=False):
    if st.button("Listar dispositivos de audio"):
        st.text_area("Dispositivos detectados", list_audio_devices(), height=280)

with st.expander("3.3 Procesar audio y generar nota clínica IA", expanded=True):
    if not st.session_state.get("encounter_id"):
        st.warning("Primero crea o selecciona una consulta.")
    else:
        use_latest = st.checkbox("Usar último audio disponible si no hay seleccionado", value=True)
        active_audio_patient_id = st.session_state.get("patient_id")
        conv_files, dict_files, other_files = list_audio_files(active_audio_patient_id)

        st.markdown("**Subir audio de la consulta**")
        upload_kind = st.radio(
            "Tipo de audio subido",
            ["conversacion", "dictado"],
            format_func=lambda x: "Conversacion" if x == "conversacion" else "Dictado medico",
            horizontal=True,
        )
        audio_file = st.file_uploader("Sube un audio de la consulta", type=AUDIO_UPLOAD_TYPES)
        if audio_file is not None:
            audio_bytes = audio_file.getvalue()
            audio_ext = os.path.splitext(audio_file.name or "")[1].lstrip(".").lower() or "wav"
            st.audio(audio_bytes, format=audio_file.type or f"audio/{audio_ext}")
            if not active_audio_patient_id:
                st.warning("Primero guarda o selecciona un paciente para asociar el audio subido.")
            else:
                uploaded_path = save_uploaded_audio(audio_file, active_audio_patient_id, upload_kind)
                if upload_kind == "dictado":
                    st.session_state["wav_dict"] = uploaded_path
                else:
                    st.session_state["wav_conv"] = uploaded_path
                st.success(f"Audio subido guardado: {audio_file_label(uploaded_path)}")

        st.markdown("**Audio disponible**")
        if active_audio_patient_id:
            st.caption(f"Mostrando solo audios del paciente ID {active_audio_patient_id}.")
        if st.session_state.get("selected_conv") not in ([""] + conv_files):
            st.session_state["selected_conv"] = ""
        if st.session_state.get("selected_dict") not in ([""] + dict_files):
            st.session_state["selected_dict"] = ""
        if st.session_state.get("selected_other") not in ([""] + other_files):
            st.session_state["selected_other"] = ""
        c1, c2 = st.columns(2)
        with c1:
            if conv_files:
                selected_conv = st.selectbox("Audio conversación", [""] + conv_files, format_func=audio_file_label, key="selected_conv")
            else:
                st.caption("No hay audios de conversación.")
                selected_conv = ""
        with c2:
            if dict_files:
                selected_dict = st.selectbox("Audio dictado", [""] + dict_files, format_func=audio_file_label, key="selected_dict")
            else:
                st.caption("No hay audios de dictado.")
                selected_dict = ""

        if other_files:
            st.selectbox("Otros archivos de audio", [""] + other_files, format_func=audio_file_label, key="selected_other")

        if st.button("Procesar y generar nota"):
            wav_conv = st.session_state.get("wav_conv")
            wav_dict = st.session_state.get("wav_dict")
            if wav_conv and not audio_belongs_to_patient(wav_conv, active_audio_patient_id):
                wav_conv = None
            if wav_dict and not audio_belongs_to_patient(wav_dict, active_audio_patient_id):
                wav_dict = None
            transcript_text, segments = "", []
            dictation_text, dictation_segments = "", []

            if selected_conv:
                wav_conv = selected_conv
            if selected_dict:
                wav_dict = selected_dict
            if use_latest:
                if not wav_conv and conv_files:
                    wav_conv = conv_files[0]
                if not wav_dict and dict_files:
                    wav_dict = dict_files[0]
                if not wav_conv and not wav_dict and other_files:
                    wav_conv = other_files[0]

            if not wav_conv and not wav_dict:
                st.error("No hay audios. Graba conversación y/o dictado.")
            else:
                if wav_conv:
                    wav_conv_name = os.path.basename(wav_conv).lower()
                    kind = "conv" if wav_conv_name.startswith("conv_") or "conversacion" in wav_conv_name else "session"
                    with st.spinner("Transcribiendo conversación con Whisper local..."):
                        transcript_text, segments = transcribe_local_with_cache(wav_conv, kind)
                    st.session_state["wav_conv"] = wav_conv
                    st.info(f"Audio conversación usado: {audio_file_label(wav_conv)}")
                if wav_dict:
                    with st.spinner("Transcribiendo dictado con Whisper local..."):
                        dictation_text, dictation_segments = transcribe_local_with_cache(wav_dict, "dict")
                    st.session_state["wav_dict"] = wav_dict
                    st.info(f"Audio dictado usado: {audio_file_label(wav_dict)}")

                st.session_state["transcript_text"] = transcript_text
                st.session_state["segments"] = segments
                st.session_state["dictation_text"] = dictation_text
                st.session_state["dictation_segments"] = dictation_segments

                if transcript_text:
                    st.markdown("**Previsualización transcripción conversación**")
                    st.text_area("", transcript_text, height=180)
                if dictation_text:
                    st.markdown("**Previsualización transcripción dictado**")
                    st.text_area("", dictation_text, height=180)

                patient = get_patient(st.session_state["patient_id"])
                encounter = get_encounter(st.session_state["encounter_id"])
                meta = {
                    "patient_name": f"{patient.get('first_name', '')} {patient.get('last_name', '')}".strip(),
                    "patient_age": "",
                    "patient_sex": patient.get("sex", ""),
                    "visit_type": encounter.get("visit_type", ""),
                    "brought_by": encounter.get("brought_by", ""),
                    "informant": encounter.get("informant", ""),
                    "chief_complaint": encounter.get("chief_complaint", ""),
                    "provider_name": encounter.get("provider_name", "") or auth_user.get("full_name", ""),
                    "facility_name": encounter.get("facility_name", ""),
                }
                img_text = (st.session_state.get("images_text") or "").strip()
                payload = {
                    "session_meta": meta,
                    "transcript_text": transcript_text,
                    "dictation_text": dictation_text,
                    "segments": segments,
                    "dictation_segments": dictation_segments,
                    "images_text": img_text,
                }
                base_key = sha1_text(json.dumps({
                    "conv_key": cache_key_for_file(wav_conv) if wav_conv else "",
                    "dict_key": cache_key_for_file(wav_dict) if wav_dict else "",
                    "schema": SCHEMA_VERSION,
                    "prompt": PROMPT_VERSION,
                    "model": LLM_MODEL,
                    "meta": meta,
                    "images_hash": sha1_text(img_text) if img_text else "",
                }, ensure_ascii=False, sort_keys=True))
                npath = note_cache_path(base_key)
                cached_note = load_json(npath)
                if cached_note:
                    note_json = cached_note
                    st.info("Nota cargada desde cache.")
                else:
                    with st.spinner("Generando nota estructurada con IA..."):
                        note_json = extract_structured(payload)
                    save_json(npath, note_json)
                    st.info("Nota guardada en cache.")

                diag_df = get_diagnoses_df(st.session_state["encounter_id"])
                full_note_txt = render_note(note_json, meta, diag_df)
                note_txt = render_soap_from_note(full_note_txt, note_json, meta)
                st.session_state["note_json"] = note_json
                st.session_state["note_txt"] = note_txt
                st.session_state["reviewed_note_encounter_id"] = st.session_state["encounter_id"]
                st.session_state["reviewed_note_text"] = note_txt
                update_encounter_content(
                    st.session_state["encounter_id"],
                    note_text=note_txt,
                    note_json=json.dumps(note_json, ensure_ascii=False),
                    transcript_text=transcript_text,
                    dictation_text=dictation_text,
                    images_text=img_text,
                    chief_complaint=fv(note_json, "chief_complaint") or encounter.get("chief_complaint", ""),
                )
                st.success("Nota clínica generada y asociada a la consulta. Revísala antes de enviarla por correo.")

# -------------------------
# 4. Diagnósticos CIE-10
# -------------------------
with st.expander("4) Diagnósticos CIE-10", expanded=True):
    if not st.session_state.get("encounter_id"):
        st.warning("Primero selecciona una consulta.")
    else:
        note_json = st.session_state.get("note_json")
        if note_json and note_json.get("suggested_icd10"):
            st.write("Sugerencias IA, deben ser validadas por el médico:")
            st.dataframe(pd.DataFrame(note_json.get("suggested_icd10")), use_container_width=True)

        term = st.text_input("Buscar CIE-10 por código o descripción", value="")
        cie_df = search_cie10(term, limit=50) if term else search_cie10("", limit=20)
        st.dataframe(cie_df, use_container_width=True, hide_index=True)

        d1, d2, d3, d4 = st.columns([1, 3, 1, 1])
        with d1:
            dx_code = st.text_input("Código", value="")
        with d2:
            dx_desc = st.text_input("Descripción", value="")
        with d3:
            dx_type = st.selectbox("Tipo", ["principal", "secundario", "antecedente", "diferencial"])
        with d4:
            dx_certainty = st.selectbox("Certeza", ["sospecha", "confirmado", "descartado"])
        if st.button("Agregar diagnóstico a la consulta"):
            if not dx_code.strip() or not dx_desc.strip():
                st.error("Debes ingresar código y descripción.")
            else:
                add_diagnosis(st.session_state["encounter_id"], dx_code, dx_desc, dx_type, dx_certainty)
                st.success("Diagnóstico agregado.")

        diag_df = get_diagnoses_df(st.session_state["encounter_id"])
        st.write("Diagnósticos registrados")
        st.dataframe(diag_df, use_container_width=True, hide_index=True)
        del_id = st.number_input("ID de diagnóstico a eliminar", min_value=0, step=1)
        if st.button("Eliminar diagnóstico seleccionado") and del_id:
            delete_diagnosis(int(del_id))
            st.success("Diagnóstico eliminado.")

# -------------------------
# 5. Nota clínica y exportaciones
# -------------------------
with st.expander("5) Nota clínica", expanded=True):
    if st.session_state.get("encounter_id"):
        encounter = get_encounter(st.session_state["encounter_id"])
        if not st.session_state.get("note_json") and encounter.get("note_json"):
            try:
                st.session_state["note_json"] = json.loads(encounter.get("note_json") or "{}")
            except Exception:
                st.session_state["note_json"] = None
        current_note = st.session_state.get("note_txt") or encounter.get("note_text") or ""
        if st.session_state.get("reviewed_note_encounter_id") != st.session_state["encounter_id"]:
            st.session_state["reviewed_note_encounter_id"] = st.session_state["encounter_id"]
            st.session_state["reviewed_note_text"] = current_note
        edited_note = st.text_area("Nota clínica editable", height=360, key="reviewed_note_text")
        if is_physician_reviewed(encounter):
            st.success("Estado clinico: Revisado por profesional")
        else:
            st.warning("Estado clinico: Borrador IA pendiente de revision profesional")
        st.caption(f"Al guardar la revisión se pedirá confirmación antes de enviar a: {(auth_user.get('email') or 'sin correo configurado').strip()}")
        if st.button("Guardar revisión y preparar envío"):
            st.session_state["note_txt"] = edited_note
            update_encounter_content(st.session_state["encounter_id"], note_text=edited_note, status="reviewed")
            with db() as conn:
                audit_event(
                    conn,
                    "note_physician_reviewed",
                    "encounter",
                    str(st.session_state["encounter_id"]),
                    user_id=auth_user.get("id"),
                    patient_id=st.session_state.get("patient_id"),
                    encounter_id=st.session_state.get("encounter_id"),
                    metadata={"source": "streamlit_note_review"},
                )
                conn.commit()
            st.session_state["pending_note_send"] = {
                "encounter_id": st.session_state["encounter_id"],
                "note_text": edited_note,
                "to_email": (auth_user.get("email") or "").strip(),
            }
            st.rerun()

        if st.session_state.get("pending_note_send"):
            confirm_reviewed_note_send_dialog()

        note_json = st.session_state.get("note_json")
        if note_json:
            st.subheader("Preguntas faltantes y contradicciones")
            st.write("Preguntas para completar:")
            st.write(note_json.get("missing_questions", []))
            st.write("Contradicciones:")
            st.write(note_json.get("contradictions", []))

        e1, e2, e3 = st.columns(3)
        with e1:
            if st.button("Exportar nota .txt"):
                out = export_txt(edited_note, f"nota_encuentro_{st.session_state['encounter_id']}_{int(time.time())}.txt")
                with db() as conn:
                    audit_event(conn, "txt_exported", "encounter", str(st.session_state["encounter_id"]), user_id=auth_user.get("id"), patient_id=st.session_state.get("patient_id"), encounter_id=st.session_state.get("encounter_id"), metadata={"path": out})
                    conn.commit()
                st.success(f"Guardado: {out}")
        with e2:
            if st.button("Exportar nota .docx"):
                out = export_docx("Ficha clínica ambulatoria", edited_note, f"nota_encuentro_{st.session_state['encounter_id']}_{int(time.time())}.docx")
                with db() as conn:
                    audit_event(conn, "docx_exported", "encounter", str(st.session_state["encounter_id"]), user_id=auth_user.get("id"), patient_id=st.session_state.get("patient_id"), encounter_id=st.session_state.get("encounter_id"), metadata={"path": out})
                    conn.commit()
                st.success(f"Guardado: {out}")
        with e3:
            if st.button("Exportar FHIR JSON"):
                patient = get_patient(st.session_state["patient_id"])
                diag_df = get_diagnoses_df(st.session_state["encounter_id"])
                out = export_fhir_bundle_json(patient, encounter, diag_df, edited_note, note_json, auth_user)
                with db() as conn:
                    audit_event(conn, "fhir_exported", "encounter", str(st.session_state["encounter_id"]), user_id=auth_user.get("id"), patient_id=st.session_state.get("patient_id"), encounter_id=st.session_state.get("encounter_id"), metadata={"path": out})
                    conn.commit()
                st.success(f"FHIR Bundle guardado: {out}")
                st.caption("Exportacion local solamente. El envio a HIS/RCE requiere FHIR_ENABLED=true, validacion medica y API formal.")
    else:
        st.caption("Aún no hay consulta activa.")

# -------------------------
# 6. HL7
# -------------------------
with st.expander("6) Integración HIS mediante mensajería HL7 v2", expanded=True):
    hl7_settings = HL7Settings.from_env()
    st.caption("Arquitectura HL7 v2.5 lista para motor externo. Por defecto no envia a HIS/RCE; genera cola y archivos locales.")
    st.info(f"Estado integracion: {hl7_settings.transport_label} | Version: {hl7_settings.version} | Requiere revision medica: {hl7_settings.require_review}")
    if not st.session_state.get("patient_id") or not st.session_state.get("encounter_id"):
        st.warning("Selecciona paciente y consulta.")
    else:
        patient = get_patient(st.session_state["patient_id"])
        encounter = get_encounter(st.session_state["encounter_id"])
        diag_df = get_diagnoses_df(st.session_state["encounter_id"])
        reviewed = is_physician_reviewed(encounter)
        st.write(f"Estado nota: {'Revisado por profesional' if reviewed else 'Borrador IA'}")
        latest_message = latest_hl7_queue_message(st.session_state["encounter_id"])
        if latest_message:
            st.write(f"Último control_id: `{latest_message.get('control_id')}`")
            st.write(f"Estado cola: `{latest_message.get('status')}` | Tipo: `{latest_message.get('message_type')}`")
        else:
            st.caption("Aún no hay mensajes HL7 en cola para esta consulta.")

        selected_hl7_type = st.selectbox("Tipo mensaje", ["MDM_T02", "ADT_A04"], format_func=lambda x: "MDM^T02 nota clinica" if x == "MDM_T02" else "ADT^A04 registro ambulatorio")
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            if st.button("Generar HL7 borrador"):
                message_id = create_hl7_draft(patient, encounter, diag_df, selected_hl7_type, auth_user.get("id"))
                st.session_state["selected_hl7_message_id"] = message_id
                st.success(f"Borrador HL7 en cola: #{message_id}")
                st.rerun()
        selected_message_id = st.session_state.get("selected_hl7_message_id") or (latest_message or {}).get("id")
        with c2:
            if st.button("Validar para envío", disabled=not selected_message_id):
                ok, msg = validate_hl7_queue_message(int(selected_message_id), get_encounter(st.session_state["encounter_id"]))
                st.success(msg) if ok else st.error(msg)
                st.rerun()
        with c3:
            if st.button("Exportar archivo HL7", disabled=not selected_message_id):
                ok, msg = export_hl7_queue_message(int(selected_message_id), auth_user.get("id"))
                st.success(f"Archivo HL7 guardado: {msg}") if ok else st.error(msg)
        with c4:
            send_enabled = hl7_settings.enabled and selected_message_id
            if st.button("Enviar HL7", disabled=not send_enabled):
                ok, msg = send_hl7_queue_message(int(selected_message_id), auth_user.get("id"))
                st.success(f"Transporte HL7 completado: {prevent_phi_in_logs(msg)}") if ok else st.error(msg)

        refreshed_message = latest_hl7_queue_message(st.session_state["encounter_id"])
        if refreshed_message:
            with st.expander("Ver payload HL7 completo (contiene datos clinicos)", expanded=False):
                st.warning("El payload puede contener datos sensibles. Usar solo para pruebas autorizadas.")
                st.text_area("Payload HL7", refreshed_message["payload"].replace("\r", "\n"), height=320)

# -------------------------
# 7. Transcripción, imágenes, chat
# -------------------------
with st.expander("7) Transcripción organizada", expanded=False):
    segments = st.session_state.get("segments", [])
    dictation_segments = st.session_state.get("dictation_segments", [])
    organized_parts = []
    if segments:
        organized_parts.append(segments_to_organized_text(segments, "CONVERSACION"))
    if dictation_segments:
        organized_parts.append(segments_to_organized_text(dictation_segments, "DICTADO MEDICO"))
    organized_text = "\n\n".join([p for p in organized_parts if p.strip()]).strip()
    st.session_state["organized_text"] = organized_text
    if organized_text:
        st.text_area("Vista", organized_text, height=320)
        if st.button("Guardar transcripción organizada .txt"):
            out = export_txt(organized_text, f"transcripcion_encuentro_{st.session_state.get('encounter_id', 'x')}_{int(time.time())}.txt")
            st.success(f"Guardado: {out}")
    else:
        st.caption("Aún no hay transcripción organizada.")

with st.expander("8) Cargar imágenes de exámenes/informes", expanded=False):
    uploaded = st.file_uploader("Sube imágenes png, jpg, jpeg", type=["png", "jpg", "jpeg"], accept_multiple_files=True)
    if uploaded:
        for f in uploaded[:6]:
            st.image(f, use_container_width=True)
        if st.button("Extraer texto y resumen clínico de imágenes"):
            with st.spinner("Analizando imágenes..."):
                instruction = "Extrae texto legible y resumen clínico corto. No inventes datos. Devuelve texto plano."
                imgs = []
                for f in uploaded:
                    b = f.read()
                    imgs.append(f"data:{f.type};base64,{base64.b64encode(b).decode('utf-8')}")
                img_text = summarize_images_with_llm(imgs, instruction)
                st.session_state["images_text"] = img_text
                if st.session_state.get("encounter_id"):
                    update_encounter_content(st.session_state["encounter_id"], images_text=img_text)
                st.success("Contenido de imágenes agregado al contexto.")
                st.text_area("Texto extraído", img_text, height=240)

with st.expander("9) Chat y búsqueda sobre la consulta", expanded=False):
    q = st.text_input("Buscar texto en transcripción", value="")
    if st.button("Buscar en segmentos"):
        hits1 = local_search_segments(st.session_state.get("dictation_segments", []), q, max_hits=15)
        hits2 = local_search_segments(st.session_state.get("segments", []), q, max_hits=15)
        st.subheader("Resultados en dictado")
        if hits1:
            for s in hits1:
                st.write(f"{float(s.get('start',0.0) or 0.0):.1f}s - {float(s.get('end',0.0) or 0.0):.1f}s: {s.get('text','')}")
        else:
            st.caption("Sin resultados.")
        st.subheader("Resultados en conversación")
        if hits2:
            for s in hits2:
                st.write(f"{float(s.get('start',0.0) or 0.0):.1f}s - {float(s.get('end',0.0) or 0.0):.1f}s: {s.get('text','')}")
        else:
            st.caption("Sin resultados.")

    st.divider()
    if "chat_messages" not in st.session_state:
        st.session_state["chat_messages"] = []
    for m in st.session_state["chat_messages"]:
        with st.chat_message(m["role"]):
            st.write(m["content"])
    user_msg = st.chat_input("Pregunta sobre la consulta, por ejemplo: ¿qué dijo sobre crisis convulsivas?")
    if user_msg:
        st.session_state["chat_messages"].append({"role": "user", "content": user_msg})
        with st.chat_message("user"):
            st.write(user_msg)
        encounter_note = ""
        if st.session_state.get("encounter_id"):
            enc = get_encounter(st.session_state["encounter_id"])
            encounter_note = enc.get("note_text") or ""
        context = (
            "CONTEXTO DISPONIBLE:\n"
            "1) Nota clínica:\n" + encounter_note[:MAX_CONTEXT_ORGANIZED] + "\n\n"
            "2) Transcripción organizada:\n" + (st.session_state.get("organized_text") or "")[:MAX_CONTEXT_ORGANIZED] + "\n\n"
            "3) Texto extraído de imágenes:\n" + (st.session_state.get("images_text") or "")[:MAX_CONTEXT_IMAGES] + "\n\n"
            "REGLAS:\n- Responde solo usando el contexto.\n- Si no está, di: 'No está mencionado'.\n- Respuesta breve y clara."
        )
        with st.chat_message("assistant"):
            with st.spinner("Pensando..."):
                resp = get_openai_client().responses.create(
                    model=CHAT_MODEL,
                    temperature=0.2,
                    input=[
                        {"role": "system", "content": "Eres un asistente clínico que responde solo con evidencia del contexto."},
                        {"role": "user", "content": context},
                        {"role": "user", "content": user_msg},
                    ],
                )
                ans = (resp.output_text or "").strip()
                st.write(ans)
                st.session_state["chat_messages"].append({"role": "assistant", "content": ans})
